# -*- coding: utf-8 -*-
r"""
通用数据解析标注平台 - 本地 Web 服务
====================================

导入一个文件夹，平台自动探测：
  - 图片：目录结构 -> 推断分类树（目录即标签）
  - 文本：txt 文件 -> 目录结构 -> 推断分类
  - 标注：检索 json / txt 标注文件，尝试解析格式
图片 / 文本独立运行；右侧 AI 助手对话式引导
（有标注 -> 分析一致性；无标注 -> 询问是否自动标注）。

启动：
  cd /d F:\Sdznsh0410\label_tool
  ..\.venv\Scripts\python.exe server.py
  浏览器打开 http://127.0.0.1:8300
"""

import os
import re
import sys
import json
import time
import random
import shutil
import hashlib
import zipfile
import threading
import traceback
from pathlib import Path
from typing import Optional, List, Any
from urllib.parse import unquote

BASE_DIR = Path(__file__).resolve().parent
ROOT_DIR = BASE_DIR.parent
sys.path.insert(0, str(ROOT_DIR))

import requests
from PIL import Image as PILImage
from fastapi import FastAPI, Request
from fastapi.responses import FileResponse, JSONResponse
from fastapi.staticfiles import StaticFiles
from pydantic import BaseModel
from starlette.concurrency import run_in_threadpool

from image_audit import compress_image

# ============================================================
# 路径与配置
# ============================================================
ANNOTATION_FILE = BASE_DIR / "annotations.json"   # 旧全局标注文件（迁移后仅作兼容回退）
REPORT_DIR = Path(os.getenv("REPORT_DIR") or (ROOT_DIR / "审核报告"))
REPORT_DIR.mkdir(parents=True, exist_ok=True)   # Docker 挂载卷启动时确保存在
THUMB_DIR = BASE_DIR / ".thumbs"
CONFIG_FILE = Path(os.getenv("CONFIG_FILE") or (BASE_DIR / "config.json"))
PROJECTS_DIR = BASE_DIR / "projects"              # 项目目录：projects/<id>/project.json + annotations.json

# 运行时数据源状态：不默认加载任何数据，必须通过 /api/load 或项目激活显式导入
# text_root 恒等于 source（同一文件夹内 txt / 图片独立扫描）
STATE = {"loaded": False, "project_id": None, "source": None, "json_root": None,
         "checkpoint": None, "text_root": None, "probe": None}

API_ENDPOINT = os.getenv("PLANNER_MODEL_ENDPOINT", "http://58.251.255.50:8229/v1")
MODEL_NAME = os.getenv("PLANNER_MODEL_NAME", "qwen3-6-35b-a3b")
VLM_TIMEOUT = 120

# 服务器默认 VLM 配置（用户未自定义时兜底）；api_key 为空 = 无鉴权（内网 vLLM）
DEFAULT_VLM = {"endpoint": API_ENDPOINT, "model": MODEL_NAME, "api_key": ""}


def vlm_cfg_of(user: str) -> dict:
    """取用户级 VLM 配置（存 config.json 的 vlm_profiles，按昵称独立）；
    未自定义或字段留空时回落服务器默认。"""
    profiles = (load_json_file(CONFIG_FILE, default={}) or {}).get("vlm_profiles") or {}
    p = profiles.get(user) or {}
    return {
        "endpoint": (p.get("endpoint") or "").strip().rstrip("/") or API_ENDPOINT,
        "model": (p.get("model") or "").strip() or MODEL_NAME,
        "api_key": (p.get("api_key") or "").strip(),
        "is_default": not bool(p),
    }


def save_vlm_profile(user: str, p: Optional[dict]):
    """保存 / 清除（p=None）用户 VLM 配置。"""
    with _project_lock:   # config.json 读改写保护
        cfg = load_json_file(CONFIG_FILE, default={}) or {}
        profiles = cfg.setdefault("vlm_profiles", {})
        if p is None:
            profiles.pop(user, None)
        else:
            profiles[user] = p
        atomic_write_text(CONFIG_FILE, json.dumps(cfg, ensure_ascii=False, indent=2))


def vlm_headers(cfg: dict) -> dict:
    """自定义 api_key 时附加 Bearer 鉴权头。"""
    return {"Authorization": f"Bearer {cfg['api_key']}"} if cfg.get("api_key") else {}

IMG_EXTS = {'.jpg', '.jpeg', '.png', '.bmp', '.webp', '.gif'}

app = FastAPI(title="标注与测试工具")

# RLock：写端点外层持锁后调用 save_annotations（内层再取锁），可重入避免死锁
_annotation_lock = threading.RLock()
_project_lock = threading.RLock()          # 保护 project.json 的读改写（成员注册等）
_tree_lock = threading.Lock()

# 未携带 X-User 请求头时的兜底身份（旧脚本 / curl 零感知）
LOCAL_USER = "本地用户"


def user_of(request: Request) -> str:
    """从请求头解析昵称身份；空则视为本地用户。

    前端对昵称做 percent-encode（HTTP 头仅允许 ASCII），此处还原中文。
    """
    name = unquote(request.headers.get("x-user") or "").strip()
    return name[:32] if name else LOCAL_USER


def touch_member(user: str):
    """带昵称的写操作触发：首次访问自动加入项目成员表，并刷新活跃时间。

    - 本地用户 / 未激活项目不记录
    - last_active 仅在超过 60 秒时落盘，避免高频写 project.json
    """
    if not user or user == LOCAL_USER or not STATE["project_id"]:
        return
    with _project_lock:
        meta = load_project(STATE["project_id"])
        if not meta:
            return
        now = int(time.time())
        for m in meta.setdefault("members", []):
            if m.get("name") == user:          # 同一昵称 = 同一身份（多设备共用）
                if now - (m.get("last_active") or 0) > 60:
                    m["last_active"] = now
                    save_project(meta)
                return
        meta["members"].append({"name": user, "role": "annotator",
                                "joined_at": now, "last_active": now})
        save_project(meta)


_TREE_CACHE = {"ts": 0, "data": None}
_TEXT_TREE_CACHE = {"ts": 0, "data": None}


# ============================================================
# 工具函数
# ============================================================
def require_loaded():
    """数据接口的未加载保护，返回 None 表示已加载"""
    if not STATE["loaded"]:
        return JSONResponse({"error": "未加载数据，请先在「数据源」中导入图片目录"}, status_code=409)
    return None


def to_rel(abs_path) -> Optional[str]:
    """绝对路径 -> 相对数据源根的 posix 路径"""
    try:
        return Path(abs_path).resolve().relative_to(STATE["source"]).as_posix()
    except Exception:
        return None


def source_abs(rel: str) -> Path:
    p = (STATE["source"] / rel).resolve()
    if not str(p).startswith(str(STATE["source"])):
        raise ValueError("非法路径")
    return p


def json_abs(rel: str) -> Optional[Path]:
    """图片 rel -> 对应 json 标注文件路径；未配置标注目录时返回 None"""
    if STATE["json_root"] is None:
        return None
    return (STATE["json_root"] / rel).with_suffix(".json")


def load_json_file(path: Path, default=None):
    try:
        if path.exists():
            return json.loads(path.read_text(encoding="utf-8"))
    except Exception:
        pass
    return default


# ============================================================
# 项目（共同标注底座）：projects/<id>/ 下 project.json + annotations.json
# ============================================================
def new_project_id() -> str:
    return "p_{}_{}".format(time.strftime("%Y%m%d_%H%M%S"), os.urandom(2).hex())


def project_dir(pid=None) -> Path:
    return PROJECTS_DIR / (pid or STATE["project_id"])


def annotations_path() -> Path:
    """激活项目的标注文件；无项目时退回旧全局文件（兼容）"""
    if STATE["project_id"]:
        return project_dir() / "annotations.json"
    return ANNOTATION_FILE


def load_project(pid: str) -> Optional[dict]:
    return load_json_file(PROJECTS_DIR / pid / "project.json", default=None)


def atomic_write_text(p: Path, text: str):
    """先写临时文件再原子替换（os.replace 同卷原子）。

    读端点无锁并发读同一文件，直接 write_text 存在撕裂读窗口；
    且写盘中途崩溃会留下截断文件，被下次写端点用空数据覆盖（标注丢失）。
    """
    tmp = p.with_name(p.name + ".tmp")
    tmp.write_text(text, encoding="utf-8")
    os.replace(tmp, p)


def save_project(meta: dict):
    with _project_lock:
        d = PROJECTS_DIR / meta["id"]
        d.mkdir(parents=True, exist_ok=True)
        atomic_write_text(d / "project.json",
                          json.dumps(meta, ensure_ascii=False, indent=2))


def list_projects() -> list:
    out = []
    if PROJECTS_DIR.is_dir():
        for d in sorted(PROJECTS_DIR.iterdir()):
            meta = load_json_file(d / "project.json", default=None)
            if meta:
                out.append({"id": meta["id"], "name": meta.get("name", ""),
                            "source": meta.get("source", "")})
    return out


def set_active_project(pid: str):
    cfg = load_json_file(CONFIG_FILE, default={}) or {}
    cfg["active_project"] = pid
    CONFIG_FILE.write_text(json.dumps(cfg, ensure_ascii=False, indent=2),
                           encoding="utf-8")


def get_active_project() -> Optional[dict]:
    pid = (load_json_file(CONFIG_FILE, default={}) or {}).get("active_project")
    return load_project(pid) if pid else None


def clear_source_caches():
    _TREE_CACHE["data"] = None
    _TREE_CACHE["ts"] = 0
    _TEXT_TREE_CACHE["data"] = None
    _TEXT_TREE_CACHE["ts"] = 0
    shutil.rmtree(THUMB_DIR, ignore_errors=True)


def ensure_default_project():
    """旧单会话数据（config.json + 全局 annotations.json）自动迁移为「默认项目」。

    仅在无有效激活项目且旧配置指向有效数据源时执行一次：
    - 全局 annotations.json 移入项目目录（数据完整保留，只是换位置）
    - box_annotations 旧路径记入 project.json，避免历史框标注丢失
    """
    PROJECTS_DIR.mkdir(exist_ok=True)
    cfg = load_json_file(CONFIG_FILE, default={}) or {}
    pid = cfg.get("active_project")
    if pid and load_project(pid):
        return
    src = str(cfg.get("source", "") or "")
    if not src or not Path(src).is_dir():
        return
    nid = new_project_id()
    save_project({"id": nid, "name": "默认项目", "source": src,
                  "box_annotations": cfg.get("box_annotations") or None,
                  "created_at": int(time.time()), "members": []})
    if ANNOTATION_FILE.exists():
        shutil.move(str(ANNOTATION_FILE), str(PROJECTS_DIR / nid / "annotations.json"))
    set_active_project(nid)


ensure_default_project()
# active_project 可能悬空（项目目录被外部删除）：校验失败则置 None，
# 标注读写回退到旧全局 annotations.json，避免写不存在的目录报错
_pid = (load_json_file(CONFIG_FILE, default={}) or {}).get("active_project")
STATE["project_id"] = _pid if (_pid and load_project(_pid)) else None


def load_annotations() -> dict:
    data = load_json_file(annotations_path(), default=None)
    if data is None:
        data = {"version": 1, "items": {}, "text_items": {}}
    if "items" not in data:
        data = {"version": 1, "items": data, "text_items": {}}
    data.setdefault("text_items", {})
    return data


def save_annotations(ann: dict):
    with _annotation_lock:
        atomic_write_text(annotations_path(),
                          json.dumps(ann, ensure_ascii=False, indent=1))


def scan_tree(force: bool = False) -> dict:
    """扫描数据源目录树（带 30s 缓存）"""
    with _tree_lock:
        if not force and _TREE_CACHE["data"] and time.time() - _TREE_CACHE["ts"] < 30:
            return _TREE_CACHE["data"]

        def walk(d: Path) -> dict:
            node = {
                "name": d.name,
                "path": d.relative_to(STATE["source"]).as_posix(),
                "children": [],
                "count": 0,   # 直接图片数
                "total": 0,   # 递归图片数
            }
            try:
                entries = sorted(d.iterdir(), key=lambda x: (x.is_file(), x.name))
            except Exception:
                return node
            for e in entries:
                if e.is_dir():
                    child = walk(e)
                    node["children"].append(child)
                elif e.suffix.lower() in IMG_EXTS:
                    node["count"] += 1
            node["total"] = node["count"] + sum(c["total"] for c in node["children"])
            return node

        # 根节点也统计直接放在数据源根下的图片（支持平铺目录）
        root = walk(STATE["source"])
        root["name"] = "全部数据"
        root["path"] = ""
        _TREE_CACHE["ts"] = time.time()
        _TREE_CACHE["data"] = root
        return root


def label_paths_with_images() -> List[str]:
    """所有含图片的目录路径（作为分类标签候选）"""
    result = []

    def walk(node):
        if node["total"] > 0:
            result.append(node["path"])
        for c in node["children"]:
            walk(c)

    for c in scan_tree()["children"]:
        walk(c)
    return sorted(result)


def label_path_of(rel_dir: str) -> str:
    """目录路径 -> ' > ' 分隔的标签路径"""
    return " > ".join(rel_dir.split("/"))


# ============================================================
# 数据探测：一次扫描识别 图片 / 文本 / 标注文件
# ============================================================
def probe_json_file(p: Path) -> dict:
    """尝试解析一个 json 文件，推断其格式类型"""
    try:
        data = json.loads(p.read_text(encoding="utf-8"))
    except Exception as e:
        return {"kind": "unreadable", "count": 0, "fields": [], "error": str(e)[:120]}
    if isinstance(data, dict) and data:
        vals = [v for v in list(data.values())[:5] if isinstance(v, dict)]
        # VLM 审核结果：key=图片路径, value 含 label_path / match
        if len(vals) >= 2 and all(("label_path" in v or "match" in v) for v in vals):
            fields = sorted({k for v in vals for k in v.keys()} - {"raw_error"})
            return {"kind": "vlm_checkpoint", "count": len(data), "fields": fields[:20]}
        if "dataList" in data:
            return {"kind": "box_annotation", "count": len(data.get("dataList", [])),
                    "fields": ["dataList.rectangle"]}
        return {"kind": "unknown_dict", "count": len(data),
                "fields": [str(k) for k in list(data.keys())[:20]]}
    if isinstance(data, list) and data:
        # 汇总标注结果：条目含 file 且能提取出框（扁平/嵌套 output/annotation 映射统一判定）
        if all(isinstance(x, dict) and "file" in x and _entry_boxes(x) for x in data[:5]):
            labels = []
            for x in data:
                if isinstance(x, dict):
                    for b in _entry_boxes(x):
                        if b["label"] not in labels:
                            labels.append(b["label"])
            return {"kind": "ann_results", "count": len(data),
                    "fields": [str(k) for k in list(data[0].keys())[:20]],
                    "labels": labels[:30]}
        # 含 file 但无法解析 → AI 兜底识别候选（/api/format/ai_parse）
        if all(isinstance(x, dict) and "file" in x for x in data[:5]):
            e0 = data[0]
            return {"kind": "unknown_ann", "count": len(data),
                    "fields": [str(k) for k in list(e0.keys())[:20]]}
        e0 = data[0]
        fields = [str(k) for k in list(e0.keys())[:20]] if isinstance(e0, dict) else []
        return {"kind": "unknown_list", "count": len(data), "fields": fields}
    return {"kind": "empty", "count": 0, "fields": []}


def probe_folder(folder: Path) -> dict:
    """扫描文件夹：图片统计 / txt 统计 / json 发现与格式识别，返回探测报告"""
    img_count = 0
    top_cats = {}          # 一级目录 -> 图片数
    txt_count = 0
    txt_dirs = set()
    same_name_json = 0     # 与图片同名的 json（框标注镜像）
    other_jsons = []       # 独立命名的 json（可能是 checkpoint / 汇总标注）

    for root, dirs, files in os.walk(folder):
        rel = Path(root).relative_to(folder).as_posix()
        top = rel.split("/")[0] if rel != "." else ""
        n_img = sum(1 for f in files if Path(f).suffix.lower() in IMG_EXTS)
        if n_img:
            img_count += n_img
            if top:
                top_cats[top] = top_cats.get(top, 0) + n_img
        txts = [f for f in files if f.lower().endswith(".txt")]
        if txts:
            txt_count += len(txts)
            txt_dirs.add(rel)
        for f in files:
            if not f.lower().endswith(".json"):
                continue
            stem = Path(f).stem
            has_img = any(Path(g).stem == stem and Path(g).suffix.lower() in IMG_EXTS
                          for g in files)
            if has_img:
                same_name_json += 1
            else:
                other_jsons.append(Path(root) / f)

    # 探测独立 json（限制数量防止超大目录拖慢导入）
    json_probes = []
    for p in other_jsons[:20]:
        info = probe_json_file(p)
        json_probes.append({"rel": p.relative_to(folder).as_posix(),
                            "size": p.stat().st_size, **info})

    # 角色推导
    checkpoint = None
    cp_probes = [j for j in json_probes if j["kind"] == "vlm_checkpoint"]
    if cp_probes:
        checkpoint = max(cp_probes, key=lambda j: j["count"])["rel"]
    json_root = folder if same_name_json > 0 else None
    results_jsons = [j["rel"] for j in json_probes if j.get("kind") == "ann_results"]
    unknown_anns = [j["rel"] for j in json_probes if j.get("kind") == "unknown_ann"]

    return {
        "folder": str(folder),
        "images": {"count": img_count,
                   "top_cats": sorted(top_cats.items(), key=lambda x: -x[1])[:20]},
        "texts": {"count": txt_count, "dirs": len(txt_dirs)},
        "jsons": json_probes,
        "same_name_json": same_name_json,
        "roles": {"json_root": json_root and str(json_root),
                  "checkpoint": checkpoint,
                  "results_jsons": results_jsons,
                  "unknown_anns": unknown_anns,
                  "text": txt_count > 0},
    }


# ============================================================
# 复核队列
# ============================================================
def build_review_items() -> List[dict]:
    ckpt = load_json_file(STATE["checkpoint"], default={}) if STATE["checkpoint"] else {}
    ann = load_annotations()["items"]
    items = []
    for abs_path, r in ckpt.items():
        rel = to_rel(abs_path)
        if rel is None:
            continue
        a = ann.get(rel, {})
        review = a.get("review") or {}
        it = {
            "image": rel,
            "label_path": r.get("label_path", ""),
            "match": r.get("match"),
            "confidence": round(float(r.get("confidence", 0) or 0), 3),
            "reason": r.get("reason", ""),
            "actual_content": r.get("actual_content", ""),
            "suggested_label": r.get("suggested_label", ""),
            "error": r.get("raw_error"),
            "reviewed": bool(review),
            "my_decision": review.get("decision"),
            "my_label": review.get("label_path", ""),
            "label_by": (a.get("classify") or {}).get("by", ""),   # 标注人（P2）
            "my_by": review.get("by", ""),                          # 复核人（P2）
        }
        # 优先级: 不匹配 > 错误 > 低置信匹配 > 匹配
        if it["match"] is False:
            pri = 0
        elif it["error"]:
            pri = 1
        elif it["match"] is True and it["confidence"] < 0.7:
            pri = 2
        else:
            pri = 3
        it["_pri"] = pri
        items.append(it)
    items.sort(key=lambda x: (x["_pri"], x["label_path"], x["image"]))
    for it in items:
        del it["_pri"]
    return items


# ============================================================
# API：数据源导入（单文件夹 + 自动探测）
# ============================================================
@app.get("/api/session")
def api_session(request: Request):
    user = user_of(request)
    touch_member(user)          # 首次访问自动加入成员表 / 刷新活跃时间
    cfg = load_json_file(CONFIG_FILE, default={}) or {}
    meta = get_active_project()
    src = str(STATE["source"]) if STATE["source"] \
        else (meta or {}).get("source") or cfg.get("source", "")
    # 服务重启后内存态丢失，但只要持久化的数据目录仍存在，就允许一键恢复
    resumable = (not STATE["loaded"]) and bool(src) and Path(src).is_dir()
    return {
        "user": user,
        "loaded": STATE["loaded"],
        "resumable": resumable,
        "folder": src,
        "project_id": (meta or {}).get("id", ""),
        "project_name": (meta or {}).get("name", ""),
        "projects": list_projects(),
    }


class ProjectCreate(BaseModel):
    name: str
    source: str = ""          # 可选：创建后立即绑定并解析数据源


@app.get("/api/projects")
def api_projects():
    return {"projects": list_projects(), "active": STATE["project_id"] or ""}


@app.post("/api/projects")
def api_projects_create(req: ProjectCreate):
    name = req.name.strip()
    if not name:
        return JSONResponse({"error": "请填写项目名称"}, status_code=400)
    meta = {"id": new_project_id(), "name": name, "source": "",
            "box_annotations": None, "created_at": int(time.time()), "members": []}
    raw = req.source.strip().strip('"')
    if raw:
        folder = Path(raw)
        if not folder.is_dir():
            return JSONResponse({"error": f"文件夹不存在: {raw}"}, status_code=400)
        meta["source"] = str(folder)
    save_project(meta)
    probe = _apply_load(Path(meta["source"]), meta) if raw else None
    # 加载成功后才持久化激活态，避免探测失败时 config 与内存态分歧
    set_active_project(meta["id"])
    return {"ok": True, "id": meta["id"], "name": meta["name"], "probe": probe}


@app.post("/api/projects/{pid}/activate")
def api_projects_activate(pid: str):
    meta = load_project(pid)
    if not meta:
        return JSONResponse({"error": "项目不存在"}, status_code=404)
    folder = str(meta.get("source") or "")
    probe = None
    if folder and Path(folder).is_dir():
        probe = _apply_load(Path(folder), meta)
        set_active_project(pid)          # 加载成功后才持久化
    else:
        # 项目未绑定数据源（或数据源已失效）：清空内存态，等待用户导入
        STATE.update({"loaded": False, "project_id": pid, "source": None,
                      "json_root": None, "checkpoint": None, "text_root": None,
                      "probe": None})
        clear_source_caches()
        set_active_project(pid)
    return {"ok": True, "probe": probe}


@app.post("/api/projects/{pid}/delete")
def api_projects_delete(pid: str):
    """删除项目：连同 projects/<pid>/ 下的标注记录、框标注与已上传数据一并删除（不可恢复）"""
    meta = load_project(pid)
    if not meta:
        return JSONResponse({"error": "项目不存在"}, status_code=404)
    was_active = pid == STATE["project_id"]
    shutil.rmtree(PROJECTS_DIR / pid, ignore_errors=True)
    if was_active:
        # 删除的是激活项目：清空内存态并取消激活，回到无项目空状态
        STATE.update({"loaded": False, "project_id": None, "source": None,
                      "json_root": None, "checkpoint": None, "text_root": None,
                      "probe": None})
        clear_source_caches()
        with _project_lock:
            cfg = load_json_file(CONFIG_FILE, default={}) or {}
            cfg.pop("active_project", None)
            atomic_write_text(CONFIG_FILE, json.dumps(cfg, ensure_ascii=False, indent=2))
    return {"ok": True, "name": meta.get("name", ""),
            "projects": list_projects(), "active": STATE["project_id"] or ""}


class LoadRequest(BaseModel):
    folder: str                  # 数据文件夹（唯一必选）


def parse_bbox(bbox):
    """bbox 兼容解析：[x1,y1,x2,y2] 或字符串 "[x1,y1,x2,y2]" → int 一维；非法返回 None"""
    if isinstance(bbox, str):
        try:
            bbox = json.loads(bbox)
        except Exception:
            return None
    if isinstance(bbox, (list, tuple)) and len(bbox) == 4:
        try:
            return [int(round(float(v))) for v in bbox]
        except (TypeError, ValueError):
            return None
    return None


def _entry_files(e: dict) -> list:
    """条目 file 字段 → [文件名]；兼容 "1.png" 与 ["1.png"]"""
    raw = e.get("file")
    if isinstance(raw, (list, tuple)):
        return [str(f).strip() for f in raw if str(f).strip()]
    fname = str(raw or "").strip()
    return [fname] if fname else []


def _entry_boxes(e: dict) -> list:
    """条目 → [{label, coordinates}]；兼容三种汇总标注格式（依次尝试）：
    ① 嵌套式 output 数组 [{type, bbox}]（label=type）
    ② annotation 映射 {label: bbox}（label=字典键，如 点选验证码{"点选文字验证码框": "[x,y,x,y]", "1": ...}）
    ③ 扁平式 条目自身 {type, bbox}"""
    out = e.get("output")
    if isinstance(out, list):
        boxes = []
        for b in out:
            if not isinstance(b, dict):
                continue
            c = parse_bbox(b.get("bbox"))
            if c:
                boxes.append({"label": str(b.get("type") or "").strip() or "object",
                              "coordinates": c})
        if boxes:
            return boxes
    ann = e.get("annotation")
    if isinstance(ann, dict):
        boxes = []
        for k, v in ann.items():
            c = parse_bbox(v)
            if c:
                boxes.append({"label": str(k).strip() or "object", "coordinates": c})
        if boxes:
            return boxes
    c = parse_bbox(e.get("bbox"))
    if c:
        return [{"label": str(e.get("type") or "").strip() or "object", "coordinates": c}]
    return []


def flat_coords(coords):
    """框坐标归一化为一维 [x1,y1,x2,y2]；兼容旧嵌套 [[x1,y1],[x2,y2]]；非法返回 None"""
    if not isinstance(coords, (list, tuple)):
        return None
    if len(coords) == 4 and all(isinstance(v, (int, float)) for v in coords):
        return [coords[0], coords[1], coords[2], coords[3]]
    if len(coords) == 2:
        try:
            (x1, y1), (x2, y2) = coords
            return [x1, y1, x2, y2]
        except (TypeError, ValueError):
            return None
    return None


def _write_grouped_mirrors(json_root, grouped) -> int:
    """grouped {rel图片路径: [{label, coordinates}]} → 写镜像 json；已存在跳过，返回写入数"""
    written = 0
    for fname, boxes in grouped.items():
        jf = (json_root / fname).with_suffix(".json")
        if jf.exists():
            continue      # 已导入过或用户已编辑过：保留现状，不覆盖
        jf.parent.mkdir(parents=True, exist_ok=True)
        mirror = {"filePath": fname, "info": {},
                  "updated_by": "import", "updated_at": int(time.time()),
                  "dataList": [{"id": i + 1, "shapeType": "rectangle", "label": b["label"],
                                "superLabel": "",
                                "coordinates": b["coordinates"],
                                "properties": {}} for i, b in enumerate(boxes)]}
        jf.write_text(json.dumps(mirror, ensure_ascii=False, indent=2),
                      encoding="utf-8")
        written += 1
    return written


def _merge_imported_labels(imported_labels: list):
    """数据导入标签并入 annotate_cfg：imported 整体替换、手动标签保留、旧版配置整体替换"""
    with _annotation_lock:
        ann = load_annotations()
        cfg = ann.get("annotate_cfg") or {}
        # 区分标签来源：数据导入的标签整体替换 imported_labels；
        # 用户此前手动创建的标签（不在旧 imported 里的）保留，接在导入标签之后。
        # 旧版本配置无 imported_labels 键：历史标签来源不可考（多为旧导入），
        # 全部随本次导入整体替换，避免换机/目录迁移后大量老旧标签残留
        if "imported_labels" in cfg:
            old_imported = cfg["imported_labels"] or []
            manual = [l for l in (cfg.get("labels") or []) if l not in old_imported]
        else:
            old_imported, manual = [], []
        merged = imported_labels + [l for l in manual if l not in imported_labels]
        cfg.update({"line": "image", "labels": merged,
                    "imported_labels": imported_labels,
                    "task_type": cfg.get("task_type") or "detect",
                    "format": cfg.get("format") or "yolo",
                    "start": cfg.get("start") or 1,
                    "group": cfg.get("group") or 50,
                    "ts": int(time.time())})
        ann["annotate_cfg"] = cfg
        save_annotations(ann)


def _apply_load(folder: Path, meta: dict) -> dict:
    """解析并挂载数据源到内存状态（导入 / 项目激活共用流程）"""
    probe = probe_folder(folder)
    checkpoint = (folder / probe["roles"]["checkpoint"]) if probe["roles"]["checkpoint"] else None
    json_root = folder if probe["roles"]["json_root"] else None

    # 无同名 json 镜像时，恢复该项目此前自动创建的独立框标注目录
    if not json_root:
        saved_jr = meta.get("box_annotations")
        if saved_jr and Path(saved_jr).is_dir():
            json_root = Path(saved_jr)

    # 汇总标注结果（每目录一个 results json：[{file, type, bbox}]）：
    # 框导入项目框标注目录（点开图片即可见可改），标签集用数据内标签自动刷新（覆盖旧缓存）
    imported_labels: List[str] = []
    if probe["roles"].get("results_jsons"):
        if not json_root:
            json_root = PROJECTS_DIR / meta["id"] / "box_annotations"
            meta["box_annotations"] = str(json_root)
            save_project(meta)
        for rel_json in probe["roles"]["results_jsons"]:
            try:
                entries = json.loads((folder / rel_json).read_text(encoding="utf-8"))
            except Exception:
                continue
            if not isinstance(entries, list):
                continue
            d = Path(rel_json).parent.as_posix()
            prefix = "" if d in (".", "") else d + "/"
            # 统一条目解析：扁平 {file,type,bbox} / 嵌套 {file:[...], output:[...]} /
            # annotation 映射 {file, annotation:{label: bbox}}（_entry_files/_entry_boxes）
            grouped: dict = {}
            for e in entries:
                if not isinstance(e, dict):
                    continue
                files = _entry_files(e)
                parsed = _entry_boxes(e)
                if not files or not parsed:
                    continue
                for b in parsed:
                    if b["label"] not in imported_labels:
                        imported_labels.append(b["label"])
                for fname in files:
                    grouped.setdefault(prefix + fname, []).extend(parsed)
            _write_grouped_mirrors(json_root, grouped)
        if imported_labels:
            _merge_imported_labels(imported_labels)

    STATE.update({"loaded": True, "project_id": meta["id"], "source": folder,
                  "json_root": json_root, "checkpoint": checkpoint,
                  "text_root": folder, "probe": probe})
    # 重置目录树缓存与缩略图缓存
    clear_source_caches()

    tree = scan_tree(force=True)
    probe["images"]["count"] = max(probe["images"]["count"], tree["total"])
    return probe


@app.post("/api/load")
def api_load(req: LoadRequest):
    raw = req.folder.strip().strip('"')
    if not raw:
        return JSONResponse({"error": "请选择数据文件夹"}, status_code=400)
    folder = Path(raw)
    if not folder.exists() or not folder.is_dir():
        return JSONResponse({"error": f"文件夹不存在: {folder}"}, status_code=400)

    # 无激活项目（首次使用）→ 自动创建「默认项目」绑定该数据源
    meta = get_active_project()
    if meta is None:
        meta = {"id": new_project_id(), "name": "默认项目", "source": str(folder),
                "box_annotations": None, "created_at": int(time.time()), "members": []}
        save_project(meta)
        _activate_after_load = True
    elif meta.get("source") != str(folder):
        # 当前项目重新导入不同文件夹：更新项目绑定的数据源
        meta["source"] = str(folder)
        save_project(meta)
        _activate_after_load = False
    else:
        _activate_after_load = False

    probe = _apply_load(folder, meta)
    if _activate_after_load:
        set_active_project(meta["id"])
    return {"ok": True, "probe": probe}


# ============================================================
# 未知标注格式 AI 兜底解析：切片 → VLM 推断「适配器/脚本」→ 试解析预览 → 用户确认后导入
# 适配器（结构化描述）优先、脚本降级；按格式指纹缓存（同结构文件复用，不再调模型）
# ============================================================

def _resolve_path(obj, path):
    """点路径取值：'file' / 'file.0' / 'output'；逐段下钻，失败返回 None"""
    cur = obj
    for part in str(path or "").split("."):
        if not part:
            continue
        if isinstance(cur, list) and part.lstrip("-").isdigit():
            try:
                cur = cur[int(part)]
            except Exception:
                return None
        elif isinstance(cur, dict):
            cur = cur.get(part)
        else:
            return None
        if cur is None:
            return None
    return cur


def _skeleton(o, depth=0):
    """去值留键结构骨架（格式指纹用：同结构 = 同格式）。
    叶子 dict 且键数 >3 视为「映射式数据」（如 annotation:{标签:坐标}，键逐图变化）→ 归一为 {map}"""
    if depth > 4:
        return "*"
    if isinstance(o, dict):
        if len(o) > 3 and all(not isinstance(v, (dict, list)) for v in o.values()):
            return "{map}"
        return "{" + ",".join(f"{k}:{_skeleton(v, depth + 1)}"
                              for k, v in list(o.items())[:10]) + "}"
    if isinstance(o, list):
        return "[" + (_skeleton(o[0], depth + 1) if o else "") + "]"
    return "*"


def _format_fingerprint(data):
    import hashlib
    return hashlib.md5(_skeleton(data[:3]).encode("utf-8")).hexdigest()[:16]


def _adapters_path():
    d = (PROJECTS_DIR / STATE["project_id"]) if STATE.get("project_id") else BASE_DIR
    d.mkdir(parents=True, exist_ok=True)
    return d / "format_adapters.json"


def _load_adapters():
    try:
        return json.loads(_adapters_path().read_text(encoding="utf-8"))
    except Exception:
        return {}


def _save_adapters(adapters):
    atomic_write_text(_adapters_path(),
                      json.dumps(adapters, ensure_ascii=False, indent=2))


def _strip_code_fence(text):
    return re.sub(r"^```[a-zA-Z]*\s*|\s*```$", "", (text or "").strip()).strip()


def _extract_json_obj(text):
    t = _strip_code_fence(text)
    try:
        return json.loads(t)
    except Exception:
        pass
    s = t.find("{")
    if s < 0:
        return None
    depth = 0
    for i in range(s, len(t)):
        if t[i] == "{":
            depth += 1
        elif t[i] == "}":
            depth -= 1
            if depth == 0:
                try:
                    return json.loads(t[s:i + 1])
                except Exception:
                    return None
    return None


def _vlm_chat(cfg, prompt, max_tokens=900):
    """纯文本 VLM 对话，返回 (content, err)"""
    payload = {"model": cfg["model"],
               "messages": [{"role": "system", "content": VLM_SYSTEM},
                            {"role": "user", "content": f"/no_think\n{prompt}"}],
               "max_tokens": max_tokens, "temperature": 0.1,
               "chat_template_kwargs": {"enable_thinking": False}}
    resp = requests.post(f"{cfg['endpoint']}/chat/completions", json=payload,
                         headers=vlm_headers(cfg), timeout=VLM_TIMEOUT)
    if resp.status_code != 200:
        return None, f"HTTP {resp.status_code}: {resp.text[:150]}"
    content = resp.json()["choices"][0]["message"]["content"]
    if "</think>" in content:
        content = content.split("</think>")[-1].strip()
    return content, None


def _run_parse_script(code, data):
    """AST 白名单校验后受限执行模型生成的 parse(data)（仅 json/re，5 秒超时）"""
    import ast
    try:
        tree = ast.parse(code)
    except SyntaxError as ex:
        return False, f"脚本语法错误: {ex}"
    allow_mods = {"json", "re"}
    forbid_names = {"open", "eval", "exec", "compile", "__import__", "globals",
                    "locals", "vars", "getattr", "setattr", "delattr", "input", "breakpoint"}
    for node in ast.walk(tree):
        if isinstance(node, ast.Import):
            for a in node.names:
                if a.name.split(".")[0] not in allow_mods:
                    return False, f"脚本禁止 import: {a.name}"
        elif isinstance(node, ast.ImportFrom):
            if (node.module or "").split(".")[0] not in allow_mods:
                return False, f"脚本禁止 import: {node.module}"
        elif isinstance(node, ast.Name) and node.id in forbid_names:
            return False, f"脚本禁止使用: {node.id}"
        elif isinstance(node, ast.Attribute) and node.attr.startswith("__"):
            return False, "脚本禁止访问双下划线属性"
    ns = {"json": json, "re": re}
    try:
        exec(compile(tree, "<adapter_script>", "exec"), ns)  # noqa: S102 - AST 白名单限制后的托管执行
    except Exception as ex:
        return False, f"脚本加载失败: {str(ex)[:150]}"
    fn = ns.get("parse")
    if not callable(fn):
        return False, "脚本未定义 parse(data) 函数"
    out = {}
    t = threading.Thread(target=lambda: out.setdefault("r", fn(data)), daemon=True)
    t.start()
    t.join(5)
    if t.is_alive():
        return False, "脚本执行超时（5 秒）"
    if "r" not in out:
        return False, "脚本无返回"
    return True, out["r"]


def _parse_with_adapter(data, adapter, script):
    """按适配器（结构化描述）或脚本解析 → (grouped{relfile:[{label,coordinates}]}, err)"""
    grouped = {}
    try:
        if adapter:
            fp_ = adapter.get("file_path") or "file"
            bp_ = adapter.get("boxes_path") or ""
            lk_ = adapter.get("label_key") or "type"
            bk_ = adapter.get("bbox_key") or "bbox"
            for e in data:
                if not isinstance(e, dict):
                    continue
                raw_f = _resolve_path(e, fp_)
                if isinstance(raw_f, (list, tuple)):
                    files = [str(f).strip() for f in raw_f if str(f).strip()]
                else:
                    fs = str(raw_f or "").strip()
                    files = [fs] if fs else []
                raw_b = _resolve_path(e, bp_)
                boxes = []
                if isinstance(raw_b, list):
                    for b in raw_b:
                        if not isinstance(b, dict):
                            continue
                        c = parse_bbox(b.get(bk_))
                        if not c:
                            continue
                        if adapter.get("bbox_type") == "xywh":
                            c = [c[0], c[1], c[0] + c[2], c[1] + c[3]]
                        boxes.append({"label": str(b.get(lk_) or "").strip() or "object",
                                      "coordinates": c})
                elif isinstance(raw_b, dict):
                    for k, v in raw_b.items():
                        c = parse_bbox(v)
                        if not c:
                            continue
                        if adapter.get("bbox_type") == "xywh":
                            c = [c[0], c[1], c[0] + c[2], c[1] + c[3]]
                        boxes.append({"label": str(k).strip() or "object",
                                      "coordinates": c})
                if files and boxes:
                    for f in files:
                        grouped.setdefault(f, []).extend(boxes)
        elif script:
            ok, result = _run_parse_script(script, data)
            if not ok:
                return {}, str(result)
            for item in result or []:
                if not isinstance(item, dict):
                    continue
                fname = str(item.get("file") or "").strip()
                boxes = []
                for b in item.get("boxes") or []:
                    try:
                        lab, coords = b[0], parse_bbox(b[1])
                    except Exception:
                        continue
                    if coords:
                        boxes.append({"label": str(lab).strip() or "object",
                                      "coordinates": coords})
                if fname and boxes:
                    grouped.setdefault(fname, []).extend(boxes)
        return grouped, ""
    except Exception as ex:
        return {}, str(ex)[:200]


class FormatPathRequest(BaseModel):
    path: str                   # 相对数据源的 json 路径


@app.post("/api/format/ai_parse")
def api_format_ai_parse(req: FormatPathRequest, request: Request):
    """未知标注格式识别：命中缓存适配器直接用，否则切片问 VLM；试解析成功才缓存并返回预览"""
    err = require_loaded()
    if err:
        return err
    p = source_abs(req.path)
    if p is None or not p.exists():
        return {"error": "文件不存在"}
    try:
        data = json.loads(p.read_text(encoding="utf-8"))
    except Exception as ex:
        return {"error": f"JSON 读取失败: {str(ex)[:100]}"}
    if not isinstance(data, list) or not data:
        return {"error": "仅支持列表型汇总标注文件"}
    sample = json.dumps(data[:5], ensure_ascii=False)[:8000]
    fp = _format_fingerprint(data)
    adapters = _load_adapters()
    cached = adapters.get(fp)
    adapter = (cached or {}).get("adapter")
    script = (cached or {}).get("script")
    src = "缓存适配器（同结构格式已识别过）"
    if not cached:
        cfg = vlm_cfg_of(user_of(request))
        prompt = f"""分析以下标注 JSON 样本（前几条），推断如何提取检测框标注。只输出一个 JSON 对象，禁止输出任何其他内容：
{{"file_path": "条目内图片文件名的点路径，如 file 或 file.0", "boxes_path": "框数据的点路径（值是框列表或{{标签:坐标}}映射），如 output 或 annotation", "label_key": "框对象内标签字段名，映射式填 null", "bbox_key": "框对象内坐标字段名，映射式填 null", "bbox_type": "auto 或 xywh（坐标为[x,y,w,h]时）"}}

样本：
{sample}"""
        content, cerr = _vlm_chat(cfg, prompt)
        if cerr:
            return {"error": f"VLM 调用失败: {cerr}"}
        adapter = _extract_json_obj(content)
        if not adapter or not adapter.get("boxes_path"):
            adapter = None
        if adapter:
            src = "AI 结构化描述"
        else:
            # 描述失败 → 降级要脚本
            prompt2 = f"""以下标注 JSON 无法用字段路径描述解析。请写 Python 函数 parse(data)（data 为完整列表），返回
[{{"file": "1.png", "boxes": [["标签", [x1,y1,x2,y2]], ...]}}, ...]
只输出代码，不要解释；仅可 import json 和 re；必须定义函数 parse。

样本：
{sample}"""
            code, cerr2 = _vlm_chat(cfg, prompt2, max_tokens=1500)
            if cerr2:
                return {"error": f"VLM 调用失败: {cerr2}"}
            code = _strip_code_fence(code)
            ok, trial = _run_parse_script(code, data[:5]) if code else (False, "空脚本")
            if not ok:
                return {"error": f"AI 脚本试解析失败: {trial}"}
            script = code
            src = "AI 生成脚本"
    grouped, perr = _parse_with_adapter(data[:5], adapter, script)
    if not grouped:
        return {"error": f"试解析失败: {perr or '未提取到任何框'}"}
    adapters[fp] = {"adapter": adapter, "script": script, "ts": int(time.time())}
    _save_adapters(adapters)
    preview = [{"file": k, "boxes": len(v)} for k, v in list(grouped.items())[:5]]
    labels = []
    for v in grouped.values():
        for b in v:
            if b["label"] not in labels:
                labels.append(b["label"])
    return {"ok": True, "source": src, "entries": len(data),
            "preview": preview, "labels": labels[:20]}


@app.post("/api/format/apply")
def api_format_apply(req: FormatPathRequest):
    """用已缓存的适配器/脚本正式解析导入：镜像幂等落盘 + 导入标签合并"""
    err = require_loaded()
    if err:
        return err
    p = source_abs(req.path)
    if p is None or not p.exists():
        return {"error": "文件不存在"}
    try:
        data = json.loads(p.read_text(encoding="utf-8"))
    except Exception as ex:
        return {"error": f"JSON 读取失败: {str(ex)[:100]}"}
    fp = _format_fingerprint(data)
    cached = _load_adapters().get(fp)
    if not cached:
        return {"error": "该文件尚未识别格式，请先「AI 识别格式」"}
    grouped, perr = _parse_with_adapter(data, cached.get("adapter"), cached.get("script"))
    if not grouped:
        return {"error": f"解析失败: {perr or '未提取到框'}"}
    # 框文件名是相对标注文件所在目录的 → 加目录前缀得到相对数据源路径
    d = Path(req.path).parent.as_posix()
    prefix = "" if d in (".", "") else d + "/"
    grouped = {prefix + k: v for k, v in grouped.items()}
    if STATE["json_root"] is None:
        if STATE["project_id"]:
            STATE["json_root"] = PROJECTS_DIR / STATE["project_id"] / "box_annotations"
        else:
            STATE["json_root"] = BASE_DIR / "box_annotations"
        STATE["json_root"].mkdir(parents=True, exist_ok=True)
        meta = get_active_project()
        if meta:
            meta["box_annotations"] = str(STATE["json_root"])
            save_project(meta)
    written = _write_grouped_mirrors(STATE["json_root"], grouped)
    labels = []
    for v in grouped.values():
        for b in v:
            if b["label"] not in labels:
                labels.append(b["label"])
    if labels:
        _merge_imported_labels(labels)
    clear_source_caches()
    scan_tree(force=True)
    return {"ok": True, "written": written, "skipped": len(grouped) - written,
            "labels": labels[:20]}


# ============================================================
# 本地数据上传（服务器部署：浏览器直传文件夹 / zip）
# 数据落在 projects/<id>/data/（Docker 挂载卷内，天然持久化）；
# 暂存目录 .staging_<sid>/ 完成后原地改名 / 解压，最终 source 指向 data/<数据集名>
# ============================================================
_upload_sessions: dict = {}          # sid -> {"dir", "kind", "name", "pid", "files", "bytes"}
_upload_lock = threading.Lock()


def _safe_upload_name(name: str) -> str:
    """数据集名 -> 合法目录名（去文件系统非法字符，限长）"""
    name = re.sub(r'[\\/:*?"<>|\r\n\t]+', "_", name or "").strip().strip(".")
    return name[:60] or "uploaded_data"


def _safe_upload_rel(rel: str) -> Optional[Path]:
    """上传相对路径 -> 安全相对 Path（去 .. / 盘符 / 绝对路径，防目录穿越）"""
    parts = []
    for p in str(rel or "").replace("\\", "/").split("/"):
        if p in ("", ".", "..") or re.fullmatch(r"[A-Za-z]:", p):
            continue
        parts.append(p)
    return Path(*parts) if parts else None


def _is_upload_junk(rel: str) -> bool:
    """系统杂项文件判定（mac / windows 资源文件），上传链路统一过滤。

    - __MACOSX 目录（mac 右键压缩 zip 必带，不滤则顶层解包失效）
    - .DS_Store（mac 目录元数据）/ desktop.ini、Thumbs.db（windows）
    - ._xxx AppleDouble 资源文件（mac 压缩包/网络盘常见，与真实文件同名成对出现）
    """
    for p in str(rel or "").replace("\\", "/").split("/"):
        low = p.lower()
        if low in ("__macosx", ".ds_store", "thumbs.db", "desktop.ini"):
            return True
        if low.startswith("._") and low != ".":
            return True
    return False


def _project_data_dir(meta: dict) -> Path:
    return PROJECTS_DIR / meta["id"] / "data"


def _cleanup_stale_uploads(data_dir: Path, max_age: int = 86400):
    """清扫放弃上传留下的孤儿暂存/回滚目录（>24h，正常流程由 cancel/finish 删除）"""
    if not data_dir.is_dir():
        return
    now = time.time()
    for d in data_dir.iterdir():
        if d.name.startswith((".staging_", ".old_")) and d.is_dir():
            try:
                if now - d.stat().st_mtime > max_age:
                    shutil.rmtree(d, ignore_errors=True)
            except OSError:
                pass


class UploadBeginRequest(BaseModel):
    name: str                    # 数据集名（作为数据源根目录名）
    kind: str = "folder"         # folder | zip


@app.post("/api/upload/begin")
def api_upload_begin(req: UploadBeginRequest, request: Request):
    """创建上传会话：建立暂存目录（不动旧数据，finish 成功后才替换）。

    与 /api/load 同语义：无激活项目时自动创建「默认项目」；重新上传在 finish 成功后替换数据源。
    """
    kind = req.kind if req.kind in ("folder", "zip") else "folder"
    meta = get_active_project()
    if meta is None:
        meta = {"id": new_project_id(), "name": "默认项目", "source": "",
                "box_annotations": None, "created_at": int(time.time()), "members": []}
        save_project(meta)
    data_dir = _project_data_dir(meta)
    data_dir.mkdir(parents=True, exist_ok=True)
    _cleanup_stale_uploads(data_dir)     # 顺带清扫历史放弃上传留下的孤儿目录
    sid = "{}_{}".format(time.strftime("%Y%m%d_%H%M%S"), os.urandom(3).hex())
    staging = data_dir / (".staging_" + sid)
    staging.mkdir(parents=True, exist_ok=True)
    with _upload_lock:
        _upload_sessions[sid] = {"dir": staging, "kind": kind,
                                 "name": _safe_upload_name(req.name), "pid": meta["id"],
                                 "files": 0, "bytes": 0}
    touch_member(user_of(request))
    return {"ok": True, "sid": sid}


@app.post("/api/upload/file")
async def api_upload_file(request: Request, sid: str, relpath: str):
    """写入一个文件：请求体即原始文件流。

    攒批 4MB 后移入线程池落盘，避免大文件上传期间同步写阻塞事件循环。
    """
    with _upload_lock:
        sess = _upload_sessions.get(sid)
    if not sess:
        return JSONResponse({"error": "上传会话不存在或已结束"}, status_code=404)
    segs = str(relpath or "").replace("\\", "/").split("/")
    if segs[0] == "" or any(s in ("", ".", "..") or re.fullmatch(r"[A-Za-z]:", s) for s in segs):
        # 显式拒绝：静默剥离 .. 会让文件落到意外位置
        return JSONResponse({"error": f"非法路径: {relpath}"}, status_code=400)
    rel = _safe_upload_rel(relpath)
    if rel is None:
        return JSONResponse({"error": f"非法路径: {relpath}"}, status_code=400)
    if _is_upload_junk(relpath):
        # mac / windows 系统杂项：不落盘，幂等返回成功（兼容未过滤杂项的旧前端，不报错重试）
        return {"ok": True, "skipped": True, "size": 0}
    dest = sess["dir"] / rel
    try:
        dest.resolve().relative_to(sess["dir"].resolve())
    except ValueError:
        return JSONResponse({"error": "非法路径"}, status_code=400)
    dest.parent.mkdir(parents=True, exist_ok=True)

    st = {"parts": [], "size": 0, "total": 0}   # parts 攒批，减少线程切换次数

    def _flush():
        mode = "wb" if st["total"] == 0 else "ab"
        with open(dest, mode) as f:
            for p in st["parts"]:
                f.write(p)
        st["total"] += st["size"]
        st["parts"] = []
        st["size"] = 0

    try:
        async for chunk in request.stream():
            st["parts"].append(chunk)
            st["size"] += len(chunk)
            if st["size"] >= 4 * 1024 * 1024:
                await run_in_threadpool(_flush)
        if st["size"] or st["total"] == 0:
            await run_in_threadpool(_flush)     # 空文件也确保落盘
    except Exception:
        dest.unlink(missing_ok=True)            # 半截文件不进数据源
        raise
    with _upload_lock:
        sess = _upload_sessions.get(sid)
        if sess:
            sess["files"] += 1
            sess["bytes"] += st["total"]
    return {"ok": True, "size": st["total"]}


@app.post("/api/upload/cancel")
def api_upload_cancel(sid: str):
    """取消上传：丢弃暂存目录"""
    with _upload_lock:
        sess = _upload_sessions.pop(sid, None)
    if sess:
        shutil.rmtree(sess["dir"], ignore_errors=True)
    return {"ok": True}


def _extract_zip(zip_path: Path, dest: Path) -> int:
    """安全解压：跳过目录项、路径穿越条目与 mac/win 系统杂项，返回解压文件数"""
    count = 0
    with zipfile.ZipFile(zip_path) as zf:
        for info in zf.infolist():
            rel = _safe_upload_rel(info.filename)
            if rel is None or info.is_dir() or _is_upload_junk(info.filename):
                continue
            target = dest / rel
            target.parent.mkdir(parents=True, exist_ok=True)
            with zf.open(info) as src, open(target, "wb") as f:
                shutil.copyfileobj(src, f)
            count += 1
    return count


@app.post("/api/upload/finish")
def api_upload_finish(request: Request, sid: str):
    """完成上传：文件夹改名 / zip 解压，复用项目激活流程解析挂载并返回探测报告。

    替换语义在此处才生效：旧数据集先改名保全，新数据解析成功后才删除，失败自动回滚。
    """
    with _upload_lock:
        sess = _upload_sessions.pop(sid, None)
    if not sess:
        return JSONResponse({"error": "上传会话不存在或已结束"}, status_code=404)
    meta = load_project(sess["pid"])
    if not meta:
        return JSONResponse({"error": "所属项目已不存在"}, status_code=404)
    staging: Path = sess["dir"]
    final = _project_data_dir(meta) / sess["name"]

    if sess["kind"] == "zip":
        zips = [p for p in staging.glob("*") if p.is_file()]   # 暂存区只会有一个 zip
        if not zips:
            shutil.rmtree(staging, ignore_errors=True)
            return JSONResponse({"error": "未收到 zip 文件"}, status_code=400)
        if zips[0].suffix.lower() != ".zip":
            shutil.rmtree(staging, ignore_errors=True)
            return JSONResponse({"error": "上传的文件不是 zip 压缩包"}, status_code=400)

    # 旧数据集改名保全（.old_* 由成功路径删除 / 失败路径回滚 / 孤儿清扫兜底）
    replaced = None
    if final.exists():
        replaced = final.parent / (final.name + ".old_" + os.urandom(2).hex())
        final.rename(replaced)

    def _rollback():
        if final.exists():
            shutil.rmtree(final, ignore_errors=True)
        if replaced and replaced.exists():
            replaced.rename(final)

    try:
        if sess["kind"] == "zip":
            final.mkdir(parents=True, exist_ok=True)
            _extract_zip(zips[0], final)
            shutil.rmtree(staging, ignore_errors=True)
            # 解包单一顶层目录（常见压缩习惯 data.zip/data/... → 直接以 data 为数据源）
            entries = list(final.iterdir())
            if len(entries) == 1 and entries[0].is_dir():
                tmp = final.parent / (final.name + ".unwrap")
                entries[0].rename(tmp)
                final.rmdir()
                tmp.rename(final)
        else:
            staging.rename(final)                    # 同卷改名，瞬间完成

        meta["source"] = str(final)
        save_project(meta)
        probe = _apply_load(final, meta)
    except zipfile.BadZipFile:
        _rollback()
        shutil.rmtree(staging, ignore_errors=True)
        return JSONResponse({"error": "zip 文件损坏或格式不支持"}, status_code=400)
    except Exception:
        _rollback()                                  # 解析失败不丢旧数据
        raise
    if replaced:
        shutil.rmtree(replaced, ignore_errors=True)
    # 单数据源语义：清掉 data/ 下其它历史数据集（新数据已探测成功接管 source）。
    # 点前缀目录（其它进行中会话的 .staging_* / .old_* / .unwrap）不动，由各自流程或孤儿清扫处理。
    for d in final.parent.iterdir():
        if d != final and d.is_dir() and not d.name.startswith("."):
            shutil.rmtree(d, ignore_errors=True)
    set_active_project(meta["id"])
    touch_member(user_of(request))
    return {"ok": True, "probe": probe,
            "files": sess["files"], "bytes": sess["bytes"]}


@app.get("/api/probe")
def api_probe():
    err = require_loaded()
    if err:
        return err
    return {"probe": STATE["probe"]}


@app.get("/api/fs/dirs")
def api_fs_dirs(path: str = ""):
    """本地目录浏览（供前端选择图片/标注目录）"""
    if not path:
        import string
        drives = [f"{d}:\\" for d in string.ascii_uppercase if Path(f"{d}:\\").exists()]
        return {"path": "", "parent": "", "dirs": drives}
    try:
        p = Path(path).resolve()
    except Exception:
        return {"error": "路径无效"}
    if not p.exists() or not p.is_dir():
        return {"error": f"目录不存在: {p}"}
    dirs = []
    try:
        for e in sorted(p.iterdir(), key=lambda x: x.name.lower()):
            if e.is_dir() and not e.name.startswith((".", "$")):
                dirs.append(str(e))
    except PermissionError:
        pass
    parent = str(p.parent) if str(p.parent) != str(p) else ""
    return {"path": str(p), "parent": parent, "dirs": dirs}


# ============================================================
# API：基础
# ============================================================
@app.get("/api/tree")
def api_tree(refresh: int = 0):
    err = require_loaded()
    if err:
        return err
    return scan_tree(force=bool(refresh))


_box_ann_cache = {"sig": None, "count": 0}


def _count_box_annotated() -> int:
    """框标注覆盖数（json_root 下 dataList 非空且对应图片存在于数据源）

    带签名缓存：先只 stat 遍历（文件数 + 最新 mtime），签名未变则直接返回缓存，
    避免每次打开总览页全量读取/解析镜像 json；写盘操作必然更新 mtime，缓存随之失效。
    图片存在性用 iterdir 精确匹配（不用 glob 模式，文件名含 [ * ? 时不误匹配/漏计）。
    """
    if not STATE["json_root"]:
        return 0
    jr = Path(STATE["json_root"])
    files, latest = [], 0.0
    for jf in jr.rglob("*.json"):
        try:
            m = jf.stat().st_mtime
        except OSError:
            continue
        files.append((jf, m))
        if m > latest:
            latest = m
    sig = (str(jr), len(files), latest)
    if _box_ann_cache["sig"] == sig:
        return _box_ann_cache["count"]
    count = 0
    for jf, _ in files:
        data = load_json_file(jf, default={})
        if not (isinstance(data, dict) and data.get("dataList")):
            continue
        stem = jf.relative_to(jr).with_suffix("")
        src_dir = STATE["source"] / stem.parent
        if src_dir.is_dir() and any(
                g.stem == stem.name and g.suffix.lower() in IMG_EXTS for g in src_dir.iterdir()):
            count += 1
    _box_ann_cache.update(sig=sig, count=count)
    return count


@app.get("/api/stats")
def api_stats():
    err = require_loaded()
    if err:
        return err
    tree = scan_tree()
    ckpt = load_json_file(STATE["checkpoint"], default={}) if STATE["checkpoint"] else {}
    ann = load_annotations()["items"]

    dist = {"total": len(ckpt), "match": 0, "mismatch": 0, "error": 0}
    for r in ckpt.values():
        if r.get("raw_error"):
            dist["error"] += 1
        elif r.get("match") is True:
            dist["match"] += 1
        elif r.get("match") is False:
            dist["mismatch"] += 1

    review_stats = {"reviewed": 0, "confirm": 0, "override": 0, "uncertain": 0}
    classify_count = 0
    recent = []
    for rel, a in ann.items():
        # 只统计当前数据源内实际存在的图片，切换数据集后旧记录不计入
        if not (STATE["source"] / rel).exists():
            continue
        rv = a.get("review")
        if rv:
            review_stats["reviewed"] += 1
            review_stats[rv.get("decision", "confirm")] = review_stats.get(rv.get("decision", "confirm"), 0) + 1
            recent.append({"image": rel, "decision": rv.get("decision"),
                           "label": rv.get("label_path", ""), "ts": rv.get("ts", 0)})
        if a.get("classify"):
            classify_count += 1
    recent.sort(key=lambda x: -x["ts"])

    # 一级分类分布
    top_cats = [{"name": c["name"], "count": c["total"]} for c in tree["children"]]

    # 框标注覆盖：json_root 下 dataList 非空、且对应图片存在于当前数据源（总览页进度用）
    box_annotated = _count_box_annotated()

    return {
        "image_total": tree["total"],
        "dir_count": len(label_paths_with_images()),
        "top_cats": top_cats,
        "vlm": dist,
        "review": review_stats,
        "classified": classify_count,
        "box_annotated": box_annotated,
        "recent": recent[:20],
    }


@app.get("/api/image")
def api_image(path: str, thumb: int = 0):
    err = require_loaded()
    if err:
        return err
    try:
        p = source_abs(path)
    except ValueError:
        return JSONResponse({"error": "非法路径"}, status_code=400)
    if not p.exists():
        return JSONResponse({"error": "文件不存在"}, status_code=404)

    if thumb:
        THUMB_DIR.mkdir(exist_ok=True)
        key = hashlib.md5(path.encode("utf-8")).hexdigest() + ".jpg"
        cache = THUMB_DIR / key
        if not cache.exists():
            img = PILImage.open(p)
            img.thumbnail((320, 320))
            if img.mode != "RGB":
                img = img.convert("RGB")
            img.save(cache, format="JPEG", quality=80)
        return FileResponse(cache, media_type="image/jpeg")

    media = {"jpg": "image/jpeg", "jpeg": "image/jpeg", "png": "image/png",
             "bmp": "image/bmp", "webp": "image/webp", "gif": "image/gif"}
    return FileResponse(p, media_type=media.get(p.suffix.lstrip(".").lower(), "application/octet-stream"))


# ============================================================
# API：复核台
# ============================================================
@app.get("/api/review/queue")
def api_review_queue(
    status: str = "pending",       # pending / reviewed / all
    result: str = "all",           # all / mismatch / error / low / match
    q: str = "",
    limit: int = 2000,
):
    err = require_loaded()
    if err:
        return err
    items = build_review_items()
    counts = {
        "total": len(items),
        "pending": sum(1 for i in items if not i["reviewed"]),
        "mismatch": sum(1 for i in items if i["match"] is False),
        "error": sum(1 for i in items if i["error"]),
        "low": sum(1 for i in items if i["match"] is True and i["confidence"] < 0.7),
    }
    filtered = []
    for it in items:
        if status == "pending" and it["reviewed"]:
            continue
        if status == "reviewed" and not it["reviewed"]:
            continue
        if result == "mismatch" and it["match"] is not False:
            continue
        if result == "error" and not it["error"]:
            continue
        if result == "low" and not (it["match"] is True and it["confidence"] < 0.7):
            continue
        if result == "match" and it["match"] is not True:
            continue
        if q and q not in it["label_path"] and q not in it["image"]:
            continue
        filtered.append(it)
    return {"items": filtered[:limit], "counts": counts}


class ReviewDecision(BaseModel):
    image: str
    decision: str            # confirm / override / uncertain
    label_path: str = ""     # override 时的正确标签（" > " 分隔）
    note: str = ""


@app.post("/api/review/decide")
def api_review_decide(d: ReviewDecision, request: Request):
    err = require_loaded()
    if err:
        return err
    user = user_of(request)
    with _annotation_lock:      # 读-改-写原子
        ann = load_annotations()
        item = ann["items"].setdefault(d.image, {})
        item["review"] = {
            "decision": d.decision,
            "label_path": d.label_path,
            "note": d.note,
            "by": user,
            "ts": int(time.time()),
        }
        save_annotations(ann)
    touch_member(user)
    return {"ok": True}


@app.post("/api/review/reset")
def api_review_reset(image: str, request: Request):
    user = user_of(request)
    with _annotation_lock:
        ann = load_annotations()
        if image in ann["items"] and "review" in ann["items"][image]:
            del ann["items"][image]["review"]
            save_annotations(ann)
    touch_member(user)
    return {"ok": True}


# ============================================================
# API：标注台
# ============================================================
@app.get("/api/label/images")
def api_label_images(dir: str, offset: int = 0, limit: int = 60,
                     unlabeled_only: int = 0):
    err = require_loaded()
    if err:
        return err
    d = STATE["source"] / dir
    if not d.exists() or not d.is_dir():
        return JSONResponse({"error": "目录不存在"}, status_code=404)
    imgs = []
    for root, _dirs, files in os.walk(d):
        for fn in files:
            if os.path.splitext(fn)[1].lower() in IMG_EXTS:
                imgs.append(Path(root) / fn)
    imgs.sort()
    ann = load_annotations()["items"]

    def img_item(f):
        rel = f.relative_to(STATE["source"]).as_posix()
        jf = json_abs(rel)
        a = ann.get(rel, {})
        cl = a.get("classify") or {}
        return {
            "image": rel,
            "name": f.name,
            "has_boxes": jf.exists() if jf else False,
            "classified": cl.get("label_path", ""),
            "classify_status": cl.get("status", ""),
            "vlm_suggest": cl.get("vlm_suggest", ""),
            "classify_by": cl.get("by", ""),       # 分类标注归属人（P2）
            "boxes_by": a.get("boxes_by", ""),     # 框标注归属人（P2）
            "reviewed": bool(a.get("review")),
        }

    if unlabeled_only:
        imgs = [f for f in imgs
                if not ((ann.get(f.relative_to(STATE["source"]).as_posix(), {})
                         or {}).get("classify") or {}).get("label_path")]
    items = [img_item(f) for f in imgs[offset:offset + limit]]
    return {"items": items, "total": len(imgs), "offset": offset, "limit": limit}


@app.get("/api/boxes")
def api_boxes(image: str):
    err = require_loaded()
    if err:
        return err
    jf = json_abs(image)
    if jf is None or not jf.exists():
        return {"exists": False, "shapes": [], "filePath": "", "updated_by": ""}
    data = load_json_file(jf, default={})
    shapes = []
    for s in data.get("dataList", []):
        if s.get("shapeType") == "point":
            # 点标注：coordinates = [x, y]
            c = s.get("coordinates", [])
            if isinstance(c, (list, tuple)) and len(c) == 2:
                try:
                    shapes.append({
                        "id": s.get("id"),
                        "label": s.get("label", ""),
                        "shapeType": "point",
                        "coordinates": [int(round(float(c[0]))), int(round(float(c[1])))],
                    })
                except (TypeError, ValueError):
                    continue
            continue
        coords = flat_coords(s.get("coordinates", []))
        if s.get("shapeType") == "line" and coords:
            # 线标注：coordinates = [x1, y1, x2, y2]（端点原序，不归一化）
            shapes.append({
                "id": s.get("id"),
                "label": s.get("label", ""),
                "shapeType": "line",
                "coordinates": coords,
            })
            continue
        if s.get("shapeType") == "rectangle" and coords:
            shapes.append({
                "id": s.get("id"),
                "label": s.get("label", ""),
                "shapeType": "rectangle",
                "coordinates": coords,
            })
    return {"exists": True, "shapes": shapes,
            "filePath": data.get("filePath", ""),
            "updated_by": data.get("updated_by", ""),   # 标注归属人（P2）
            "size": data.get("info", {})}


class BoxShape(BaseModel):
    label: str
    shapeType: Optional[str] = None   # rectangle / point / line；缺省时按坐标形态推断（旧客户端兼容）
    coordinates: List[Any]    # 框 [x1,y1,x2,y2]（一维；兼容旧客户端嵌套 [[x1,y1],[x2,y2]]）；点 [x,y]；线 [x1,y1,x2,y2]


class BoxSaveRequest(BaseModel):
    image: str
    shapes: List[BoxShape]


@app.post("/api/boxes")
def api_boxes_save(req: BoxSaveRequest, request: Request):
    err = require_loaded()
    if err:
        return err
    user = user_of(request)
    # 读图片尺寸放锁外：纯读固有属性，不参与互斥；避免大图 IO 阻塞全局写锁
    img_path = source_abs(req.image)
    if not img_path.exists():
        return JSONResponse({"error": "图片不存在"}, status_code=404)
    try:
        with PILImage.open(img_path) as im:
            w, h = im.size
    except Exception:
        return JSONResponse({"error": "图片读取失败"}, status_code=400)
    with _annotation_lock:      # 镜像 json 与 annotations.json 的读改写原子
        if STATE["json_root"] is None:
            # 数据源无同名 json 镜像：自动启用项目目录存放框标注（不污染数据源）
            if STATE["project_id"]:
                STATE["json_root"] = PROJECTS_DIR / STATE["project_id"] / "box_annotations"
            else:
                STATE["json_root"] = BASE_DIR / "box_annotations"
            STATE["json_root"].mkdir(parents=True, exist_ok=True)
            meta = get_active_project()
            if meta:
                meta["box_annotations"] = str(STATE["json_root"])
                save_project(meta)

        jf = json_abs(req.image)
        jf.parent.mkdir(parents=True, exist_ok=True)

        old = load_json_file(jf, default=None) if jf.exists() else None
        # 首次编辑前备份
        if jf.exists():
            bak = jf.with_name(jf.name + ".bak")
            if not bak.exists():
                shutil.copy2(jf, bak)

        data_list = []
        for i, s in enumerate(req.shapes, 1):
            raw = s.coordinates
            st = (s.shapeType or "").strip().lower()
            is_pair = isinstance(raw, (list, tuple)) and len(raw) == 2
            # 点标注：显式 point + [x, y]；或未声明 shapeType 的旧客户端发长度 2 纯标量（长度 2 嵌套是旧客户端框）
            if (st == "point" and is_pair) or (not st and is_pair
                    and all(isinstance(v, (int, float)) and not isinstance(v, bool) for v in raw)):
                try:
                    data_list.append({
                        "id": i,
                        "shapeType": "point",
                        "label": s.label,
                        "superLabel": "",
                        "coordinates": [int(round(float(raw[0]))), int(round(float(raw[1])))],
                        "properties": {},
                    })
                    continue
                except (TypeError, ValueError):
                    continue
            # 线标注：显式 line + 长度 4 的纯标量 [x1, y1, x2, y2]（端点原序，不 min/max）
            if st == "line" and isinstance(raw, (list, tuple)) and len(raw) == 4 \
                    and all(isinstance(v, (int, float)) and not isinstance(v, bool) for v in raw):
                try:
                    data_list.append({
                        "id": i,
                        "shapeType": "line",
                        "label": s.label,
                        "superLabel": "",
                        "coordinates": [int(round(float(v))) for v in raw],
                        "properties": {},
                    })
                    continue
                except (TypeError, ValueError):
                    continue
            coords = flat_coords(raw)
            if not coords:
                continue
            x1 = min(int(coords[0]), int(coords[2])); x2 = max(int(coords[0]), int(coords[2]))
            y1 = min(int(coords[1]), int(coords[3])); y2 = max(int(coords[1]), int(coords[3]))
            data_list.append({
                "id": i,
                "shapeType": "rectangle",
                "label": s.label,
                "superLabel": "",
                "coordinates": [x1, y1, x2, y2],
                "properties": {},
            })

        out = {
            "filePath": (old or {}).get("filePath", req.image),
            "info": {"height": h, "width": w, "depth": 3},
            "groupsProperties": (old or {}).get("groupsProperties", []),
            "qaMap": (old or {}).get("qaMap", {}),
            "updated_by": user,       # 标注归属（P2）
            "updated_at": int(time.time()),
            "dataList": data_list,
        }
        jf.write_text(json.dumps(out, ensure_ascii=False, indent=2), encoding="utf-8")

        # annotations.json 同步记录框标注归属人
        ann = load_annotations()
        ann["items"].setdefault(req.image, {})["boxes_by"] = user
        save_annotations(ann)
    touch_member(user)
    return {"ok": True, "saved": len(data_list), "file": str(jf)}


class ClassifyRequest(BaseModel):
    image: str
    label_path: str          # 标签体系路径（" > " 分隔）
    vlm_suggest: str = ""
    note: str = ""


@app.post("/api/classify")
def api_classify(req: ClassifyRequest, request: Request):
    err = require_loaded()
    if err:
        return err
    user = user_of(request)
    with _annotation_lock:      # 读-改-写原子
        ann = load_annotations()
        item = ann["items"].setdefault(req.image, {})
        # label_path 为空 → 清除该图的分类标注（取消预标注 / 手工标注）
        if not req.label_path.strip():
            item.pop("classify", None)
            save_annotations(ann)
            return {"ok": True, "cleared": True}
        old = item.get("classify") or {}
        item["classify"] = {
            "label_path": req.label_path,
            "vlm_suggest": req.vlm_suggest or old.get("vlm_suggest", ""),
            "status": "confirmed",
            "note": req.note,
            "by": user,           # 分类标注归属人（P2）
            "ts": int(time.time()),
        }
        save_annotations(ann)
    touch_member(user)
    return {"ok": True}


# ============================================================
# API：数据集预览 + 标注配置 + 标注结果导出
# ============================================================
@app.get("/api/dataset/preview")
def api_dataset_preview(sample: int = 300):
    """数据集概览：图片格式分布 / 尺寸抽样 / 体积统计（预览模式观察用）"""
    err = require_loaded()
    if err:
        return err
    formats, total_bytes, files = {}, 0, []
    for root, _, fns in os.walk(STATE["source"]):
        for fn in fns:
            ext = os.path.splitext(fn)[1].lower()
            if ext not in IMG_EXTS:
                continue
            p = Path(root) / fn
            try:
                sz = p.stat().st_size
            except OSError:
                continue
            total_bytes += sz
            key = ext.lstrip(".")
            formats[key] = formats.get(key, 0) + 1
            files.append((str(p.relative_to(STATE["source"])).replace("\\", "/"), sz))
    files.sort()

    # 抽样读取尺寸（只读 header，不解码像素）
    sizes = {"sampled": 0, "min_w": 0, "max_w": 0, "avg_w": 0, "min_h": 0, "max_h": 0, "avg_h": 0}
    sw = sh = 0
    for rel, _ in files[:max(1, min(sample, len(files)))]:
        try:
            with PILImage.open(source_abs(rel)) as im:
                w, h = im.size
        except Exception:
            continue
        n = sizes["sampled"]
        sizes["min_w"] = w if n == 0 else min(sizes["min_w"], w)
        sizes["max_w"] = max(sizes["max_w"], w)
        sizes["min_h"] = h if n == 0 else min(sizes["min_h"], h)
        sizes["max_h"] = max(sizes["max_h"], h)
        sw += w; sh += h
        sizes["sampled"] = n + 1
    if sizes["sampled"]:
        sizes["avg_w"] = round(sw / sizes["sampled"])
        sizes["avg_h"] = round(sh / sizes["sampled"])

    probe = STATE.get("probe") or {}
    txts = (probe.get("texts") or {})
    return {
        "images": {
            "count": len(files),
            "total_bytes": total_bytes,
            "formats": sorted(formats.items(), key=lambda x: -x[1]),
            "sizes": sizes,
        },
        "texts": {"count": txts.get("count", 0), "dirs": txts.get("dirs", 0)},
    }


class AnnotateConfigRequest(BaseModel):
    line: str = "image"            # image / text
    task_type: str = "classify"    # classify / detect（文本线固定 classify）
    format: str = "txt"            # classify: txt/csv；detect: yolo/coco/json
    start: int = 1                 # 起标点（从第几张/条开始，1-based）
    group: int = 50                # 分组数量
    labels: List[str] = []         # 标签集


@app.get("/api/annotate/config")
def api_get_annotate_config():
    """读取上次保存的标注配置"""
    err = require_loaded()
    if err:
        return err
    return {"ok": True, "config": load_annotations().get("annotate_cfg") or {}}


@app.post("/api/annotate/config")
def api_annotate_config(req: AnnotateConfigRequest):
    """保存标注配置（持久化，供导出与下次会话恢复）"""
    err = require_loaded()
    if err:
        return err
    with _annotation_lock:
        ann = load_annotations()
        old = ann.get("annotate_cfg") or {}
        labels = [l.strip() for l in req.labels if l.strip()]
        # 维护标签来源：导入标签中被用户删除/改名的自动移出 imported_labels；新增的视为手动创建
        kept_imported = [l for l in (old.get("imported_labels") or []) if l in labels]
        ann["annotate_cfg"] = {
            "line": req.line,
            "task_type": req.task_type,
            "format": req.format,
            "start": max(1, req.start),
            "group": max(1, req.group),
            "labels": labels,
            "imported_labels": kept_imported,
            "ts": int(time.time()),
        }
        save_annotations(ann)
    return {"ok": True, "config": ann["annotate_cfg"]}


def _collect_box_annotations():
    """收集全部框标注：[(rel, w, h, [(label, x1, y1, x2, y2)])]

    镜像 info 缺宽高（数据导入的镜像 info 为空字典）时，按 filePath
    回读数据源图片尺寸（只读 header），保证导入框也能导出 YOLO/COCO。
    """
    out = []
    if not STATE["json_root"]:
        return out
    for root, _, fns in os.walk(STATE["json_root"]):
        for fn in fns:
            if not fn.lower().endswith(".json") or fn.endswith(".bak"):
                continue
            p = Path(root) / fn
            rel = str(p.relative_to(STATE["json_root"])).replace("\\", "/")
            data = load_json_file(p, default={})
            info = data.get("info") or {}
            w, h = int(info.get("width") or 0), int(info.get("height") or 0)
            shapes = []
            for s in data.get("dataList", []):
                coords = flat_coords(s.get("coordinates", []))
                if s.get("shapeType") == "rectangle" and coords:
                    x1, y1, x2, y2 = (min(int(coords[0]), int(coords[2])), min(int(coords[1]), int(coords[3])),
                                      max(int(coords[0]), int(coords[2])), max(int(coords[1]), int(coords[3])))
                    shapes.append((s.get("label", ""), x1, y1, x2, y2))
            if shapes and not (w and h) and STATE.get("source"):
                fp = str(data.get("filePath") or "").replace("\\", "/")
                try:
                    imgp = STATE["source"] / fp
                    if fp and imgp.exists():
                        with PILImage.open(imgp) as im:
                            w, h = im.size
                except Exception:
                    pass
            if shapes and w and h:
                out.append((rel, w, h, shapes))
    return out


def _collect_point_annotations():
    """收集全部点标注：[(rel, [(label, x, y)])]；框镜像中的 shapeType=point 条目"""
    out = []
    if not STATE["json_root"]:
        return out
    for root, _, fns in os.walk(STATE["json_root"]):
        for fn in fns:
            if not fn.lower().endswith(".json") or fn.endswith(".bak"):
                continue
            p = Path(root) / fn
            rel = str(p.relative_to(STATE["json_root"])).replace("\\", "/")
            data = load_json_file(p, default={})
            pts = []
            for s in data.get("dataList", []):
                if s.get("shapeType") != "point":
                    continue
                c = s.get("coordinates", [])
                if isinstance(c, (list, tuple)) and len(c) == 2:
                    try:
                        pts.append((s.get("label", ""),
                                    int(round(float(c[0]))), int(round(float(c[1])))))
                    except (TypeError, ValueError):
                        continue
            if pts:
                out.append((rel, pts))
    return out


def _collect_line_annotations():
    """收集全部线标注：[(rel, [(label, x1, y1, x2, y2)])]；框镜像中的 shapeType=line 条目（端点原序）"""
    out = []
    if not STATE["json_root"]:
        return out
    for root, _, fns in os.walk(STATE["json_root"]):
        for fn in fns:
            if not fn.lower().endswith(".json") or fn.endswith(".bak"):
                continue
            p = Path(root) / fn
            rel = str(p.relative_to(STATE["json_root"])).replace("\\", "/")
            data = load_json_file(p, default={})
            segs = []
            for s in data.get("dataList", []):
                if s.get("shapeType") != "line":
                    continue
                c = s.get("coordinates", [])
                if isinstance(c, (list, tuple)) and len(c) == 4 \
                        and all(isinstance(v, (int, float)) and not isinstance(v, bool) for v in c):
                    segs.append((s.get("label", ""),
                                 int(round(float(c[0]))), int(round(float(c[1]))),
                                 int(round(float(c[2]))), int(round(float(c[3])))))
            if segs:
                out.append((rel, segs))
    return out


def _new_export_dir(kind: str) -> Path:
    """每次导出生成独立时间戳子目录：export/<kind>_<时间>[/序号]/

    分批多次导出互不覆盖；mkdir 原子创建（exist_ok=False + 冲突重试），
    多用户同秒并发导出也不会共用同一目录。
    """
    base = BASE_DIR / "export"
    ts = time.strftime("%Y%m%d_%H%M%S")
    d = base / f"{kind}_{ts}"
    n = 1
    while True:
        try:
            d.mkdir(parents=True, exist_ok=False)   # 原子创建：同目录并发只有一方成功
            return d
        except FileExistsError:
            n += 1
            d = base / f"{kind}_{ts}_{n}"


@app.post("/api/annotate/export")
def api_annotate_export(req: AnnotateConfigRequest):
    """按配置的格式导出标注结果到工具目录 export/（不污染数据源）

    - 每次导出独立时间戳子目录，分批导出不互相覆盖
    - YOLO 标签文件按数据源相对目录结构落盘，不同子目录同名图片（如
      click/1.png 与 slider_end/1.png）不再冲突
    """
    err = require_loaded()
    if err:
        return err
    ann = load_annotations()

    if req.task_type == "classify":
        rows = [(rel, (v.get("classify") or {}).get("label_path", ""))
                for rel, v in ann["items"].items()
                if (v.get("classify") or {}).get("label_path")]
        # 文本线：导出标注池已确认结果
        if req.line == "text":
            rows = [(e.get("text", ""), e.get("label_path", ""))
                    for e in (ann.get("text_pool") or {}).values()
                    if e.get("label_path")]
        if not rows:
            return {"error": "没有已保存的分类标注可导出"}
        out_dir = _new_export_dir(f"classify_{req.format}")
        if req.format == "csv":
            f = out_dir / "classify.csv"
            f.write_text("\ufeff" + "item,label\n" +
                         "\n".join(f'"{a.replace(chr(34), chr(34)*2)}","{b}"' for a, b in rows),
                         encoding="utf-8")
        else:
            f = out_dir / "classify.txt"
            f.write_text("\n".join(f"{a}\t{b}" for a, b in rows), encoding="utf-8")
        return {"ok": True, "count": len(rows), "file": str(f)}

    # ---- detect ----
    if not STATE["json_root"]:
        return {"error": "未检测到框标注目录，无法导出检测标注"}
    boxes = _collect_box_annotations()
    points = _collect_point_annotations()
    segs = _collect_line_annotations()
    lines_skipped = sum(len(s) for _, s in segs)   # 线暂无 YOLO/COCO 等价物，导出框时明示跳过数
    if req.format == "txt":
        # txt（检测线）= 点/线标注导出：点 每行 图片路径<TAB>标签<TAB>x<TAB>y；
        # 线 每行 图片路径<TAB>标签<TAB>x1<TAB>y1<TAB>x2<TAB>y2（框请用 YOLO/COCO）
        if not points and not segs:
            return {"error": "没有点/线标注可导出（框标注请选择 YOLO / COCO 格式）"}
        out_dir = _new_export_dir("points")
        n_rows, files = 0, []
        if points:
            plines = [f"{rel}\t{label}\t{x}\t{y}" for rel, pts in points for label, x, y in pts]
            f_pts = out_dir / "points.txt"
            f_pts.write_text("\n".join(plines), encoding="utf-8")
            n_rows += len(plines)
            files.append(f_pts)
        if segs:
            slines = [f"{rel}\t{label}\t{x1}\t{y1}\t{x2}\t{y2}"
                      for rel, ss in segs for label, x1, y1, x2, y2 in ss]
            f_segs = out_dir / "lines.txt"
            f_segs.write_text("\n".join(slines), encoding="utf-8")
            n_rows += len(slines)
            files.append(f_segs)
        return {"ok": True, "count": n_rows, "file": str(files[0] if len(files) == 1 else out_dir),
                "labels": len({lb for _, pts in points for lb, _, _ in pts}
                              | {lb for _, ss in segs for lb, *_ in ss})}
    if not boxes:
        return {"error": "没有框标注可导出"}
    label_set = sorted({s[0] for _, _, _, shapes in boxes for s in shapes})
    if req.format == "yolo":
        yolo_dir = _new_export_dir("yolo")
        (yolo_dir / "classes.txt").write_text("\n".join(label_set), encoding="utf-8")
        for rel, w, h, shapes in boxes:
            # 标签文件镜像图片相对路径（click/1.json → click/1.txt），避免同名 stem 冲突
            lab = yolo_dir / Path(rel).with_suffix(".txt")
            lab.parent.mkdir(parents=True, exist_ok=True)
            lines = []
            for label, x1, y1, x2, y2 in shapes:
                cid = label_set.index(label)
                cx, cy = (x1 + x2) / 2 / w, (y1 + y2) / 2 / h
                bw, bh = (x2 - x1) / w, (y2 - y1) / h
                lines.append(f"{cid} {cx:.6f} {cy:.6f} {bw:.6f} {bh:.6f}")
            lab.write_text("\n".join(lines), encoding="utf-8")
        return {"ok": True, "count": len(boxes), "labels": len(label_set),
                "file": str(yolo_dir),
                "points_skipped": sum(len(p) for _, p in points),
                "lines_skipped": lines_skipped}
    if req.format == "coco":
        images, annotations = [], []
        aid = 1
        for iid, (rel, w, h, shapes) in enumerate(boxes, 1):
            images.append({"id": iid, "file_name": rel, "width": w, "height": h})
            for label, x1, y1, x2, y2 in shapes:
                annotations.append({
                    "id": aid, "image_id": iid, "category_id": label_set.index(label) + 1,
                    "bbox": [x1, y1, x2 - x1, y2 - y1], "area": (x2 - x1) * (y2 - y1),
                    "iscrowd": 0,
                })
                aid += 1
        coco = {
            "info": {"description": "exported by 通用数据解析标注平台",
                     "date": time.strftime("%Y-%m-%d %H:%M:%S")},
            "categories": [{"id": i + 1, "name": n} for i, n in enumerate(label_set)],
            "images": images, "annotations": annotations,
        }
        f = _new_export_dir("coco") / "annotations_coco.json"
        f.write_text(json.dumps(coco, ensure_ascii=False, indent=1), encoding="utf-8")
        return {"ok": True, "count": len(boxes), "labels": len(label_set), "file": str(f),
                "points_skipped": sum(len(p) for _, p in points),
                "lines_skipped": lines_skipped}
    return {"error": f"不支持的检测导出格式：{req.format}"}


# ---- 图片批量 VLM 预标注（当前目录未标注图片） ----
BATCH_VLM_STATE = {"running": False, "total": 0, "done": 0, "errors": 0, "dir": "", "stop": False}


def collect_pending_images(dir: str) -> List[str]:
    """收集目录下未标注图片的相对路径（dir 为空 = 整个数据源）"""
    d = STATE["source"] / dir if dir else STATE["source"]
    if not d.is_dir():
        return []
    imgs = []
    for root, _dirs, files in os.walk(d):
        for fn in files:
            if os.path.splitext(fn)[1].lower() in IMG_EXTS:
                p = Path(root) / fn
                try:
                    imgs.append(p.relative_to(STATE["source"]).as_posix())
                except ValueError:
                    continue  # 越界路径（如传入绝对路径），跳过
    imgs.sort()
    ann = load_annotations()["items"]
    return [rel for rel in imgs
            if not ((ann.get(rel, {}) or {}).get("classify") or {}).get("label_path")]


def run_batch_vlm(images: List[str], workers: int, cfg: dict):
    from concurrent.futures import ThreadPoolExecutor, as_completed
    BATCH_VLM_STATE.update(running=True, total=len(images), done=0, errors=0, stop=False)

    def one(rel):
        try:
            if BATCH_VLM_STATE["stop"]:
                return rel, None, "stopped"
            r = call_vlm_classify(rel, cfg)
            if not r.get("ok"):
                return rel, None, r.get("error")
            return rel, r, None
        except Exception as e:
            return rel, None, f"{type(e).__name__}: {e}"

    ex = ThreadPoolExecutor(max_workers=workers)
    try:
        futures = [ex.submit(one, i) for i in images]
        for future in as_completed(futures):
            rel, r, err = future.result()
            if err == "stopped":
                continue
            if r:
                with _annotation_lock:   # 与前端写端点共用全局锁，避免并发覆盖
                    ann = load_annotations()
                    item = ann["items"].setdefault(rel, {})
                    item["classify"] = {
                        "label_path": r["label_path"],
                        "vlm_suggest": r["label_path"],
                        "vlm_confidence": r.get("confidence", 0),
                        "vlm_reason": r.get("reason", ""),
                        "status": "pending",
                        "ts": int(time.time()),
                    }
                    save_annotations(ann)
            else:
                BATCH_VLM_STATE["errors"] += 1
            BATCH_VLM_STATE["done"] += 1
            if BATCH_VLM_STATE["stop"]:
                break
    finally:
        ex.shutdown(wait=False, cancel_futures=True)
        BATCH_VLM_STATE["running"] = False


@app.post("/api/label/batch_vlm")
def api_label_batch_vlm(request: Request, dir: str, limit: int = 50, workers: int = 4):
    err = require_loaded()
    if err:
        return err
    if BATCH_VLM_STATE["running"]:
        return JSONResponse({"error": "批量预标注进行中"}, status_code=409)
    pending = collect_pending_images(dir)[:limit]
    if not pending:
        return {"ok": True, "started": 0, "message": "当前目录没有未标注图片"}
    user = user_of(request)
    touch_member(user)
    BATCH_VLM_STATE["dir"] = dir
    # 以触发者的 VLM 配置跑完本批次（用户级配置）
    cfg = vlm_cfg_of(user)
    threading.Thread(target=run_batch_vlm, args=(pending, workers, cfg), daemon=True).start()
    return {"ok": True, "started": len(pending)}


@app.get("/api/label/batch_vlm/status")
def api_label_batch_vlm_status():
    return BATCH_VLM_STATE


@app.post("/api/label/batch_vlm/stop")
def api_label_batch_vlm_stop():
    if not BATCH_VLM_STATE["running"]:
        return {"ok": True, "stopped": False, "message": "当前没有进行中的批量预标注"}
    BATCH_VLM_STATE["stop"] = True
    return {"ok": True, "stopped": True, "done": BATCH_VLM_STATE["done"], "total": BATCH_VLM_STATE["total"]}


@app.get("/api/label/batch_vlm/pending")
def api_label_batch_vlm_pending(dir: str = ""):
    """查询未标注图片数量（用于启动前确认）"""
    err = require_loaded()
    if err:
        return err
    return {"ok": True, "count": len(collect_pending_images(dir))}


class BatchClearRequest(BaseModel):
    dir: str = ""


def _vlm_pending_rels(dir: str) -> List[str]:
    """收集批量预标注生成的分类记录（status=pending 且有 vlm_suggest）"""
    prefix = (dir.strip("/") + "/") if dir.strip("/") else ""
    ann = load_annotations()["items"]
    out = []
    for rel, item in ann.items():
        if prefix and not rel.startswith(prefix):
            continue
        cls = (item.get("classify") or {})
        if cls.get("status") == "pending" and cls.get("vlm_suggest"):
            out.append(rel)
    return out


@app.get("/api/label/batch_vlm/clear_count")
def api_label_batch_vlm_clear_count(dir: str = ""):
    """查询可删除的批量预标注数量"""
    err = require_loaded()
    if err:
        return err
    return {"ok": True, "count": len(_vlm_pending_rels(dir))}


@app.post("/api/label/batch_vlm/clear")
def api_label_batch_vlm_clear(req: BatchClearRequest):
    """删除批量预标注结果（不影响手工标注与复核结果）"""
    err = require_loaded()
    if err:
        return err
    if BATCH_VLM_STATE["running"]:
        return JSONResponse({"error": "批量预标注进行中，请先停止"}, status_code=409)
    with _annotation_lock:
        ann = load_annotations()
        rels = _vlm_pending_rels(req.dir)
        for rel in rels:
            item = ann["items"].get(rel)
            if not item:
                continue
            item.pop("classify", None)
            if not item:
                ann["items"].pop(rel, None)
        save_annotations(ann)
    return {"ok": True, "removed": len(rels)}


# ============================================================
# API：VLM 预标注
# ============================================================
VLM_SYSTEM = "你是一名专业的内容安全审核专家，精通中国互联网内容审核标准。直接输出JSON结果，禁止输出分析过程或任何解释性文字。"


def call_vlm_classify(image_rel: str, cfg: dict) -> dict:
    img_path = source_abs(image_rel)
    img_b64 = compress_image(str(img_path))
    candidates = label_paths_with_images()
    if not candidates:
        return {"ok": False, "error": "图片目录中没有含图片的子目录，无法预标注"}
    cand_text = "\n".join(f"{i}. {p}" for i, p in enumerate(candidates, 1))
    prompt = f"""判断这张图片属于以下哪个分类标签路径，从候选列表中选择唯一最匹配的一项。

候选标签路径（"/" 分隔的目录层级，如 数据集/一级分类/二级分类）：
{cand_text}

直接输出JSON（禁止输出思考过程）：
{{"label_path": "原样复制选中的标签路径", "confidence": 0.0到1.0, "reason": "30字以内中文理由"}}"""

    url = f"{cfg['endpoint']}/chat/completions"
    payload = {
        "model": cfg["model"],
        "messages": [
            {"role": "system", "content": VLM_SYSTEM},
            {"role": "user", "content": [
                {"type": "image_url", "image_url": {"url": f"data:image/jpeg;base64,{img_b64}"}},
                {"type": "text", "text": f"/no_think\n{prompt}"},
            ]},
        ],
        "max_tokens": 800,
        "temperature": 0.1,
        # vLLM：chat_template_kwargs 必须在 payload 顶层（openai extra_body 语义）
        "chat_template_kwargs": {"enable_thinking": False},
    }
    resp = requests.post(url, json=payload, headers=vlm_headers(cfg), timeout=VLM_TIMEOUT)
    if resp.status_code != 200:
        return {"ok": False, "error": f"HTTP {resp.status_code}: {resp.text[:200]}"}
    content = resp.json()["choices"][0]["message"]["content"]
    if "</think>" in content:
        content = content.split("</think>")[-1].strip()

    # 解析 JSON
    result = None
    try:
        result = json.loads(content)
    except Exception:
        for m in re.finditer(r'\{[^{}]*\}', content, re.DOTALL):
            try:
                obj = json.loads(m.group(0))
                if "label_path" in obj:
                    result = obj
            except Exception:
                continue
    if not result or "label_path" not in result:
        m = re.search(r'"label_path"\s*:\s*"([^"]+)"', content)
        if m:
            result = {"label_path": m.group(1), "confidence": 0.5, "reason": "（正则提取）"}
    if not result:
        return {"ok": False, "error": f"解析失败: {content[:200]}"}

    lp = str(result.get("label_path", "")).strip()
    # 容错：模型可能输出 " > " 分隔，统一为目录 "/" 格式
    norm = lp.replace(" > ", "/").replace(">", "/").replace("\\", "/").strip("/")
    if norm in candidates:
        lp = norm
    elif lp in candidates:
        pass
    else:
        # 模糊匹配：路径尾部对齐
        cands = [c for c in candidates if c.endswith(lp) or lp.endswith(c)
                 or c.endswith(norm) or norm.endswith(c)]
        if cands:
            lp = sorted(cands, key=len)[0]
        else:
            lp = norm
    return {
        "ok": True,
        "label_path": lp,
        "confidence": float(result.get("confidence", 0.5) or 0.5),
        "reason": str(result.get("reason", ""))[:100],
    }


@app.post("/api/vlm/predict")
def api_vlm_predict(image: str, request: Request):
    err = require_loaded()
    if err:
        return err
    try:
        return call_vlm_classify(image, vlm_cfg_of(user_of(request)))
    except Exception as e:
        return {"ok": False, "error": f"{type(e).__name__}: {e}"}


@app.get("/api/vlm/status")
def api_vlm_status(request: Request):
    cfg = vlm_cfg_of(user_of(request))
    try:
        r = requests.get(f"{cfg['endpoint']}/models", headers=vlm_headers(cfg), timeout=10)
        return {"ok": r.status_code == 200, "endpoint": cfg["endpoint"],
                "model": cfg["model"], "is_default": cfg["is_default"]}
    except Exception as e:
        return {"ok": False, "endpoint": cfg["endpoint"], "model": cfg["model"],
                "is_default": cfg["is_default"], "error": str(e)}


class VlmModelsRequest(BaseModel):
    endpoint: str = ""          # 留空 = 用当前用户已保存的 endpoint
    api_key: str = ""           # 留空 = 用当前用户已保存的 key / 默认（无鉴权）


@app.post("/api/vlm/models")
def api_vlm_models(req: VlmModelsRequest, request: Request):
    """按给定 endpoint / key 拉取可用模型列表（模型设置弹窗内下拉选择模型名，保存前即可测试连通）"""
    cfg = vlm_cfg_of(user_of(request))
    endpoint = (req.endpoint or "").strip().rstrip("/") or cfg["endpoint"]
    key = (req.api_key or "").strip() or cfg["api_key"]
    if not endpoint.startswith(("http://", "https://")):
        return JSONResponse({"error": "API 地址需以 http:// 或 https:// 开头"}, status_code=400)
    try:
        r = requests.get(f"{endpoint}/models",
                         headers=vlm_headers({"api_key": key}), timeout=10)
        if r.status_code != 200:
            return {"ok": False, "error": f"HTTP {r.status_code}: {r.text[:200]}", "models": []}
        ids = sorted({m.get("id") for m in r.json().get("data", [])
                      if isinstance(m, dict) and m.get("id")})
        return {"ok": True, "endpoint": endpoint, "models": ids}
    except Exception as e:
        return {"ok": False, "error": str(e), "models": []}


# ---- 用户级 VLM 配置（界面可改，按昵称独立，默认兜底） ----
class VlmConfigRequest(BaseModel):
    endpoint: str = ""
    model: str = ""
    api_key: str = ""


@app.get("/api/vlm/config")
def api_vlm_config_get(request: Request):
    user = user_of(request)
    return {**vlm_cfg_of(user), "user": user}


@app.post("/api/vlm/config")
def api_vlm_config_set(req: VlmConfigRequest, request: Request):
    """保存当前用户的 VLM 配置；三个字段全空 = 恢复服务器默认。"""
    user = user_of(request)
    p = {"endpoint": req.endpoint.strip(), "model": req.model.strip(),
         "api_key": req.api_key.strip()}
    if not any(p.values()):
        save_vlm_profile(user, None)
        return {"ok": True, "reset": True, **vlm_cfg_of(user)}
    # 校验 endpoint 格式（避免保存明显无效的地址）
    if not p["endpoint"].startswith(("http://", "https://")):
        return JSONResponse({"error": "API 地址需以 http:// 或 https:// 开头"}, status_code=400)
    save_vlm_profile(user, p)
    return {"ok": True, **vlm_cfg_of(user)}


# ============================================================
# 文本线：txt 目录扫描 / 行读取 / 导入标注池
# ============================================================
def scan_text_dirs() -> List[dict]:
    """扫描文本目录：[{rel_dir, files: [Path...]}]"""
    out = []
    for root, dirs, files in os.walk(STATE["text_root"]):
        txts = sorted(f for f in files if f.lower().endswith(".txt"))
        if not txts:
            continue
        rel_dir = Path(root).relative_to(STATE["text_root"]).as_posix()
        out.append({"rel_dir": rel_dir, "files": [Path(root) / f for f in txts]})
    return out


# ============================================================
# 文本线 API：目录树 / 行预览 / 导入标注池
# ============================================================
def scan_text_tree() -> dict:
    """扫描 txt 分布目录树（带缓存）：{name, path, count, children}"""
    with _tree_lock:
        if _TEXT_TREE_CACHE["data"] and time.time() - _TEXT_TREE_CACHE["ts"] < 30:
            return _TEXT_TREE_CACHE["data"]
        dirs = {}
        for d in scan_text_dirs():
            n = len(d["files"])
            parts = [p for p in d["rel_dir"].split("/") if p and p != "."]
            key = "/".join(parts)
            dirs[key] = dirs.get(key, 0) + n
        root = {"name": "全部文本", "path": "", "count": 0, "children": []}
        node_map = {"": root}
        for key in sorted(dirs):
            parts = key.split("/")
            for i in range(1, len(parts) + 1):
                pk = "/".join(parts[:i])
                if pk not in node_map:
                    parent = node_map["/".join(parts[:i - 1])]
                    node = {"name": parts[i - 1], "path": pk, "count": 0, "children": []}
                    parent["children"].append(node)
                    node_map[pk] = node
            node_map[key]["count"] = dirs[key]
        def roll(n):
            n["total"] = n["count"] + sum(roll(c) for c in n["children"])
            return n["total"]
        root["total"] = roll(root)
        _TEXT_TREE_CACHE["ts"] = time.time()
        _TEXT_TREE_CACHE["data"] = root
        return root


def collect_text_lines(dir_filter: str = "") -> List[dict]:
    """收集目录（可过滤）下所有 txt 的文本行：[{text, en, code, file, line, label_path}]"""
    rows = []
    for d in scan_text_dirs():
        if dir_filter and not (d["rel_dir"] == dir_filter
                               or d["rel_dir"].startswith(dir_filter + "/")):
            continue
        label_path = label_path_of(d["rel_dir"])
        for fp in d["files"]:
            rel_file = fp.relative_to(STATE["text_root"]).as_posix()
            try:
                with open(fp, "r", encoding="utf-8", errors="ignore") as f:
                    for i, line in enumerate(f, 1):
                        parts = line.rstrip("\r\n").split("\t")
                        if not parts or not parts[0].strip():
                            continue
                        rows.append({
                            "text": parts[0].strip(),
                            "en": parts[1].strip() if len(parts) > 1 else "",
                            "code": parts[2].strip() if len(parts) > 2 else "",
                            "file": rel_file, "line": i, "label_path": label_path,
                        })
            except Exception:
                continue
    return rows


@app.get("/api/text/tree")
def api_text_tree(refresh: int = 0):
    err = require_loaded()
    if err:
        return err
    if refresh:
        _TEXT_TREE_CACHE["data"] = None
    return scan_text_tree()


@app.get("/api/text/lines")
def api_text_lines(dir: str = "", offset: int = 0, limit: int = 50):
    err = require_loaded()
    if err:
        return err
    rows = collect_text_lines(dir.strip())
    return {"items": rows[offset:offset + limit], "total": len(rows),
            "offset": offset, "limit": limit}


@app.post("/api/text/import")
def api_text_import(dir: str = "", limit: int = 100000):
    """把 txt 文本行导入标注池（去重），供 LLM 预标注 -> 人工确认"""
    err = require_loaded()
    if err:
        return err
    rows = collect_text_lines(dir.strip())[:limit]
    with _annotation_lock:
        ann = load_annotations()
        pool = ann.setdefault("text_pool", {})
        added = dup = 0
        for r in rows:
            tid = text_pool_id(r["text"])
            if tid in pool:
                dup += 1
                continue
            pool[tid] = {"text": r["text"], "label_path": "", "en": r["en"],
                         "code": r["code"], "src_label": r["label_path"],
                         "vlm": None, "status": "pending", "ts": int(time.time())}
            added += 1
        save_annotations(ann)
    return {"ok": True, "added": added, "duplicated": dup, "total": len(pool)}


# ============================================================
# 文本标注池：未标注文本进来 -> LLM 预标注 -> 人工确认 -> 导出
# ============================================================
class TextPoolAdd(BaseModel):
    texts: List[str]


POOL_VLM_STATE = {"running": False, "total": 0, "done": 0, "errors": 0}


def text_label_paths() -> List[str]:
    """文本标注候选标签：文本目录结构优先，回退图片目录结构（' > ' 格式）"""
    if not STATE["loaded"]:        # 服务重启后未加载数据时，目录树不可用（避免 None 崩溃）
        return []
    paths = []
    if STATE.get("text_root"):
        for d in scan_text_dirs():
            if d["rel_dir"]:
                paths.append(label_path_of(d["rel_dir"]))
    if not paths:
        paths = [label_path_of(p) for p in label_paths_with_images() if p]
    return sorted(set(p for p in paths if p))


def call_text_classify(text: str, cfg: dict) -> dict:
    """LLM 给单条未标注文本判标签（从数据目录结构选；cfg 为触发者的 VLM 配置）"""
    candidates = text_label_paths()
    if not candidates:
        return {"ok": False, "error": "没有可用标签（图片目录无含图子目录且未配置文本目录）"}
    cand_text = "\n".join(f"{i}. {p}" for i, p in enumerate(candidates, 1))
    prompt = f"""判断以下文本内容属于哪个分类标签路径，从候选列表中选择唯一最匹配的一项。如果都不匹配，选择最接近的一项。

候选标签路径（" > " 分隔的多级标签）：
{cand_text}

待分类文本：
{text[:2000]}

直接输出JSON（禁止输出思考过程）：
{{"label_path": "原样复制选中的标签路径", "confidence": 0.0到1.0, "reason": "30字以内中文理由"}}"""

    payload = {
        "model": cfg["model"],
        "messages": [
            {"role": "system", "content": VLM_SYSTEM},
            {"role": "user", "content": f"/no_think\n{prompt}"},
        ],
        "max_tokens": 800,
        "temperature": 0.1,
        "chat_template_kwargs": {"enable_thinking": False},
    }
    resp = requests.post(f"{cfg['endpoint']}/chat/completions", json=payload,
                         headers=vlm_headers(cfg), timeout=VLM_TIMEOUT)
    if resp.status_code != 200:
        return {"ok": False, "error": f"HTTP {resp.status_code}: {resp.text[:200]}"}
    content = resp.json()["choices"][0]["message"]["content"]
    if "</think>" in content:
        content = content.split("</think>")[-1].strip()

    result = None
    try:
        result = json.loads(content)
    except Exception:
        # 提取最后一个完整 JSON 对象
        for m in re.finditer(r'\{[^{}]*\}', content, re.DOTALL):
            try:
                obj = json.loads(m.group(0))
                if "label_path" in obj:
                    result = obj
            except Exception:
                continue
    if not result or not result.get("label_path"):
        m = re.search(r'"label_path"\s*:\s*"([^"]+)"', content)
        if m:
            result = {"label_path": m.group(1)}
    if not result or not result.get("label_path"):
        return {"ok": False, "error": f"解析失败: {content[:150]}"}

    lp = str(result["label_path"]).strip()
    norm = lp.replace("/", " > ").replace("\\", " > ")
    norm = " > ".join(x.strip() for x in norm.split(">") if x.strip())
    if norm in candidates:
        lp = norm
    elif lp not in candidates:
        cands = [c for c in candidates if c.endswith(norm) or norm.endswith(c)]
        lp = sorted(cands, key=len)[0] if cands else norm
    return {"ok": True, "label_path": lp,
            "confidence": float(result.get("confidence", 0.5) or 0.5),
            "reason": str(result.get("reason", ""))[:100]}


def text_pool_id(text: str) -> str:
    return hashlib.md5(text.strip().encode("utf-8")).hexdigest()[:16]


@app.post("/api/text_pool/add")
def api_text_pool_add(req: TextPoolAdd):
    with _annotation_lock:
        ann = load_annotations()
        pool = ann.setdefault("text_pool", {})
        added, dup = 0, 0
        for t in req.texts:
            t = (t or "").strip()
            if not t:
                continue
            tid = text_pool_id(t)
            if tid in pool:
                dup += 1
                continue
            pool[tid] = {"text": t, "label_path": "", "en": "", "code": "",
                         "vlm": None, "status": "pending", "ts": int(time.time())}
            added += 1
        save_annotations(ann)
    return {"ok": True, "added": added, "duplicated": dup}


@app.get("/api/text_pool/list")
def api_text_pool_list(status: str = "all", q: str = ""):
    ann = load_annotations()
    pool = ann.get("text_pool", {})
    items = []
    for tid, e in pool.items():
        if status == "pending" and e.get("label_path"):
            continue
        if status == "labeled" and not e.get("label_path"):
            continue
        if status == "confirmed" and e.get("status") != "confirmed":
            continue
        if q and q not in e.get("text", "") and q not in e.get("label_path", ""):
            continue
        items.append({"id": tid, **e})
    items.sort(key=lambda x: -x.get("ts", 0))
    return {"items": items, "total": len(pool),
            "pending": sum(1 for e in pool.values() if not e.get("label_path")),
            "confirmed": sum(1 for e in pool.values() if e.get("status") == "confirmed"),
            "labels": text_label_paths()}


def run_pool_vlm(ids: List[str], workers: int, cfg: dict):
    from concurrent.futures import ThreadPoolExecutor, as_completed
    POOL_VLM_STATE.update(running=True, total=len(ids), done=0, errors=0)

    def one(tid):
        ann = load_annotations()
        text = ann["text_pool"].get(tid, {}).get("text", "")
        if not text:
            return tid, None, "条目不存在"
        try:
            return tid, call_text_classify(text, cfg), None
        except Exception as e:
            return tid, None, f"{type(e).__name__}: {e}"

    try:
        with ThreadPoolExecutor(max_workers=workers) as ex:
            futures = [ex.submit(one, i) for i in ids]
            for future in as_completed(futures):
                tid, r, err = future.result()
                if r and r.get("ok"):
                    with _annotation_lock:   # 与前端写端点共用全局锁，避免并发覆盖
                        ann = load_annotations()
                        e = ann["text_pool"].get(tid)
                        if e is not None:
                            e["vlm"] = {"label_path": r["label_path"],
                                        "confidence": r.get("confidence", 0),
                                        "reason": r.get("reason", "")}
                            e["label_path"] = r["label_path"]
                            e["status"] = "pending"
                            save_annotations(ann)
                else:
                    with _annotation_lock:
                        ann = load_annotations()
                        e = ann["text_pool"].get(tid)
                        if e is not None:
                            e["vlm"] = {"error": (r or {}).get("error") or err}
                            save_annotations(ann)
                    POOL_VLM_STATE["errors"] += 1
                POOL_VLM_STATE["done"] += 1
    finally:
        POOL_VLM_STATE["running"] = False


@app.post("/api/text_pool/vlm")
def api_text_pool_vlm(request: Request, workers: int = 4, limit: int = 100):
    if POOL_VLM_STATE["running"]:
        return JSONResponse({"error": "预标注进行中"}, status_code=409)
    if not text_label_paths():
        return JSONResponse({"error": "没有可用标签（图片目录无含图子目录且未配置文本目录）"}, status_code=409)
    ann = load_annotations()
    ids = [tid for tid, e in ann.get("text_pool", {}).items()
           if not e.get("label_path")][:limit]
    if not ids:
        return {"ok": True, "started": 0, "message": "没有待预标注的文本"}
    user = user_of(request)
    touch_member(user)
    # 以触发者的 VLM 配置跑完本批次（用户级配置）
    cfg = vlm_cfg_of(user)
    threading.Thread(target=run_pool_vlm, args=(ids, workers, cfg), daemon=True).start()
    return {"ok": True, "started": len(ids)}


@app.get("/api/text_pool/vlm/status")
def api_text_pool_vlm_status():
    return POOL_VLM_STATE


class PoolDecide(BaseModel):
    id: str
    label_path: str


@app.post("/api/text_pool/decide")
def api_text_pool_decide(d: PoolDecide, request: Request):
    user = user_of(request)
    with _annotation_lock:
        ann = load_annotations()
        e = ann.get("text_pool", {}).get(d.id)
        if e is None:
            return JSONResponse({"error": "条目不存在"}, status_code=404)
        if d.label_path:
            e["label_path"] = d.label_path
        e["status"] = "confirmed"
        e["by"] = user          # 文本标注归属人（P2）
        e["ts"] = int(time.time())
        save_annotations(ann)
    touch_member(user)
    return {"ok": True}


@app.post("/api/text_pool/delete")
def api_text_pool_delete(id: str):
    with _annotation_lock:
        ann = load_annotations()
        if id in ann.get("text_pool", {}):
            del ann["text_pool"][id]
            save_annotations(ann)
    return {"ok": True}


@app.post("/api/text_pool/export")
def api_text_pool_export():
    from openpyxl import Workbook
    from openpyxl.styles import Font
    from openpyxl.utils import get_column_letter
    ann = load_annotations()
    pool = ann.get("text_pool", {})
    confirmed = [e for e in pool.values() if e.get("status") == "confirmed"]

    ts = time.strftime("%Y%m%d_%H%M%S")
    excel_path = REPORT_DIR / f"文本标注结果_{ts}.xlsx"
    csv_path = REPORT_DIR / f"文本标注结果_{ts}.csv"

    wb = Workbook()
    ws = wb.active
    ws.title = "标注结果"
    headers = ["文本", "标签路径", "英文标签", "编码", "来源"]
    for col, h in enumerate(headers, 1):
        ws.cell(row=1, column=col, value=h).font = Font(name="微软雅黑", bold=True)
    for r, e in enumerate(confirmed, 2):
        src = "LLM预标注" if e.get("vlm") and e["vlm"].get("label_path") == e.get("label_path") else "人工指定"
        ws.cell(row=r, column=1, value=e["text"][:500])
        ws.cell(row=r, column=2, value=e.get("label_path", ""))
        ws.cell(row=r, column=3, value=e.get("en", ""))
        ws.cell(row=r, column=4, value=e.get("code", ""))
        ws.cell(row=r, column=5, value=src)
    for col, w in enumerate([60, 50, 22, 18, 12], 1):
        ws.column_dimensions[get_column_letter(col)].width = w
    wb.save(excel_path)

    with open(csv_path, "w", encoding="utf-8-sig", newline="") as f:
        f.write("文本,标签路径,英文标签,编码\n")
        for e in confirmed:
            cells = [e["text"][:500].replace(",", "，"), e.get("label_path", ""),
                     e.get("en", ""), e.get("code", "")]
            f.write(",".join(cells) + "\n")
    return {"ok": True, "excel": str(excel_path), "csv": str(csv_path),
            "confirmed": len(confirmed)}


# ============================================================
# AI 助手：对话式引导（右侧面板，意图路由）
# ============================================================
class AssistantRequest(BaseModel):
    message: str
    line: str = "image"                 # image / text
    history: List[dict] = []            # [{"role","content"}] 最近几轮


ASSISTANT_ACTION_DOC = """- start_batch_label: 对图片发起 VLM 批量预标注。args: {"dir": 目录筛选(相对路径,空=全部), "limit": 数量上限}
- open_review: 打开复核台查看/处理预标注结果。args: {}
- import_texts: 把指定目录的文本行导入文本标注池。args: {"dir": 目录筛选(空=全部), "limit": 数量上限}
- batch_label_texts: 对文本标注池中待标注文本发起 LLM 批量预标注。args: {"limit": 数量上限}
- export_results: 导出已确认的文本标注结果。args: {}
- none: 无需操作，仅回答用户问题"""


ASSISTANT_SYSTEM = """你是通用数据解析标注平台的 AI 助手，负责引导用户完成图片分类标注与文本标注。
当前数据状态（供你判断与回答）：
{context}

可用动作（用户确认执行时才返回对应 action）：
{actions}

严格输出 JSON（禁止输出思考过程）：
{{"reply": "给用户的简短中文回复", "action": "动作名", "args": {{}}}}

规则：
1. 一次只返回一个动作；纯问答/解释类 action 用 none
2. args 未涉及的参数用默认值（dir="", limit=200）
3. 回复要具体：引用数据状态中的真实数字
4. 数据里没有任何标注文件时，主动询问用户是否需要 AI 自动预标注
5. 用户模糊时先澄清，不要擅自执行批量操作"""


def assistant_context(line: str) -> str:
    """把当前数据状态压缩成文本上下文注入 system prompt"""
    lines = []
    probe = STATE.get("probe") or {}
    imgs = probe.get("images", {}) or {}
    txts = probe.get("texts", {}) or {}
    lines.append(f"数据根目录: {STATE['source']}")

    cats = imgs.get("top_cats") or []
    if isinstance(cats, dict):
        cat_pairs = list(cats.items())
    else:
        cat_pairs = [tuple(c[:2]) for c in cats if isinstance(c, (list, tuple)) and len(c) >= 2]
    cat_str = "; ".join(f"{k}({v}张)" for k, v in cat_pairs[:10]) or "无一级目录（可能平铺）"
    lines.append(f"图片: 共{imgs.get('count', 0)}张，一级目录: {cat_str}")

    if txts.get("count"):
        n_dirs = txts.get("dirs")
        n_dirs = n_dirs if isinstance(n_dirs, int) else len(n_dirs or [])
        lines.append(f"文本: 共{txts['count']}行txt，分布在{n_dirs}个目录")
    else:
        lines.append("文本: 无 txt 文件")

    jsons = probe.get("jsons") or []
    if jsons:
        kinds = {}
        for j in jsons:
            kinds[j.get("kind", "?")] = kinds.get(j.get("kind", "?"), 0) + 1
        lines.append(f"JSON标注文件: {kinds}")
    else:
        lines.append("JSON标注文件: 未发现")

    roles = probe.get("roles") or {}
    if roles.get("checkpoint"):
        lines.append(f"检测到 VLM 预标注结果文件: {roles['checkpoint']}")
    if roles.get("json_root"):
        lines.append("检测到与图片同名的框标注 JSON（镜像标注）")

    ann = load_annotations()
    if line == "image":
        items = ann.get("items", {})
        labeled = sum(1 for v in items.values() if (v.get("classify") or {}).get("label_path"))
        reviewed = sum(1 for v in items.values() if v.get("review", {}).get("decision"))
        lines.append(f"图片进度: 已标注{labeled}张 / 已复核{reviewed}张 / 共{imgs.get('count', 0)}张")
    else:
        pool = ann.get("text_pool", {})
        pending = sum(1 for e in pool.values() if e.get("status") == "pending")
        confirmed = sum(1 for e in pool.values() if e.get("status") == "confirmed")
        lines.append(f"文本池进度: 待标注{pending} / 已确认{confirmed} / 共{len(pool)}条")
    return "\n".join(lines)


def call_assistant_llm(context: str, history: List[dict], message: str, cfg: dict) -> dict:
    system = ASSISTANT_SYSTEM.format(context=context, actions=ASSISTANT_ACTION_DOC)
    msgs = [{"role": "system", "content": system}]
    for h in history[-6:]:
        if isinstance(h, dict) and h.get("role") in ("user", "assistant") and h.get("content"):
            msgs.append({"role": h["role"], "content": str(h["content"])[:1500]})
    msgs.append({"role": "user", "content": f"/no_think\n{message}"})

    payload = {
        "model": cfg["model"],
        "messages": msgs,
        "max_tokens": 600,
        "temperature": 0.3,
        "chat_template_kwargs": {"enable_thinking": False},
    }
    resp = requests.post(f"{cfg['endpoint']}/chat/completions", json=payload,
                         headers=vlm_headers(cfg), timeout=VLM_TIMEOUT)
    if resp.status_code != 200:
        return {"ok": False, "error": f"HTTP {resp.status_code}: {resp.text[:200]}"}
    content = resp.json()["choices"][0]["message"]["content"]
    if "</think>" in content:
        content = content.split("</think>")[-1].strip()

    result = None
    try:
        result = json.loads(content)
    except Exception:
        for m in re.finditer(r'\{.*\}', content, re.DOTALL):
            try:
                result = json.loads(m.group(0))
                break
            except Exception:
                continue
    if not isinstance(result, dict) or "reply" not in result:
        # 兜底：把原始输出当回复
        return {"ok": True, "reply": content[:500] or "（模型无输出）", "action": "none", "args": {}}

    # 兜底：模型把动作 JSON 嵌套进 reply 字符串时，解包内层
    if str(result.get("action", "none")).strip() in ("", "none") and isinstance(result.get("reply"), str):
        try:
            inner = json.loads(result["reply"])
            if isinstance(inner, dict) and inner.get("reply") and inner.get("action"):
                result = inner
        except Exception:
            pass

    action = str(result.get("action", "none")).strip()
    if action not in ("start_batch_label", "open_review", "import_texts",
                      "batch_label_texts", "export_results", "none"):
        action = "none"
    args = result.get("args") if isinstance(result.get("args"), dict) else {}
    return {"ok": True, "reply": str(result["reply"])[:800],
            "action": action, "args": args}


@app.post("/api/assistant")
def api_assistant(req: AssistantRequest, request: Request):
    err = require_loaded()
    if err:
        return err
    if not req.message.strip():
        return JSONResponse({"error": "消息为空"}, status_code=400)
    try:
        ctx = assistant_context(req.line if req.line in ("image", "text") else "image")
        return call_assistant_llm(ctx, req.history, req.message.strip(),
                                  vlm_cfg_of(user_of(request)))
    except Exception as e:
        traceback.print_exc()
        return {"ok": False, "error": f"{type(e).__name__}: {e}"}


# ============================================================
# 静态资源
# ============================================================
app.mount("/", StaticFiles(directory=str(BASE_DIR / "static"), html=True), name="static")


if __name__ == "__main__":
    import uvicorn
    print("=" * 60)
    print("标注与测试工具")
    print(f"  http://127.0.0.1:8300")
    print(f"  VLM: {MODEL_NAME} @ {API_ENDPOINT}")
    print("  数据源：未预加载，请在页面「数据源」中导入")
    print("=" * 60)
    uvicorn.run(app, host="127.0.0.1", port=8300, log_level="warning")
